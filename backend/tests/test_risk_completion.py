import copy
import json
from datetime import datetime, timedelta
from io import BytesIO
from uuid import uuid4

import pytest
import pytest_asyncio
from sqlalchemy import DateTime, DefaultClause, MetaData, select, text

from app.core.exceptions import ConflictError, ForbiddenError, ValidationError, NotFoundError
from app.models import Base
from app.models.quality_risk import ProductCategory, QualityProduct, QualityRiskCase, QualitySourceRecord
from app.models.task import InspectionTask
from app.models.user import User
from app.models.export_job import ExportJob
from app.schemas.demo_risk_workflow import DemoDraftReview
from app.schemas.quality_risk import QualityEventCreate
from app.schemas.risk_completion import LabDecision, LabReturn, RiskReportCreate
from app.schemas.user import CurrentUser
from app.services.risk_lab_service import RiskLabService
from app.services.risk_report_service import RiskReportService
from app.services.export_job_service import ExportJobService
from app.services.report_renderers import build_renderer
from test_quality_risk_v4 import domain  # noqa: F401
from test_demo_risk_workflow import setup, submitted
from test_risk_collaboration import collaboration  # noqa: F401
from test_demo_sampling import sampling  # noqa: F401


@pytest_asyncio.fixture
async def completion(domain, monkeypatch):
    db, service, wf, case, source = await setup(domain, monkeypatch)
    category, product = str(uuid4()), str(uuid4())
    db.add(ProductCategory(id=category, org_id=service("expert").org_id, code="DEMO-C", name="模拟类别", is_active=True))
    db.add(QualityProduct(id=product, org_id=service("expert").org_id, category_id=category, name="模拟产品", is_active=True))
    scope = {"execution_mode":"demo", "product_id":product,"category_id":category}
    (await db.get(QualityRiskCase,case["id"])).scope = scope
    (await db.get(QualitySourceRecord,source["id"])).product_ref = {"product_id":product,"category_id":category}
    await db.flush()
    assessment = await submitted(wf(),case,source)
    assessment = await wf("expert2").review(case["id"],assessment["id"],DemoDraftReview(expected_revision=3,decision="approve",comment="已核对模拟资料及限制"))
    metadata = MetaData()
    table = Base.metadata.tables["export_jobs"].to_metadata(metadata)
    for c in table.c:
        c.server_default = DefaultClause(text("CURRENT_TIMESTAMP")) if isinstance(c.type,DateTime) and c.server_default is not None else None
    async with db.bind.begin() as conn: await conn.run_sync(metadata.create_all)
    def lab(role="expert"):return RiskLabService(db,service(role).current)
    def reports(role="expert"):return RiskReportService(db,service(role).current)
    decision = LabDecision(assessment_id=assessment["id"],expected_basis_hash=assessment["demo_validation"]["basis_hash"],mode="skip",purpose="模拟资料分析阶段，暂不需要实验室验证")
    report = RiskReportCreate(assessment_id=assessment["id"],expected_basis_hash=decision.expected_basis_hash)
    return db, service, wf, case, source, assessment, product, category, lab, reports, decision, report


@pytest.mark.asyncio
async def test_skipping_validation_report_requires_no_task_and_is_idempotent(completion):
    db, _, _, case, _, assessment, _, _, lab, reports, decision, payload = completion
    # A report also works without making any laboratory decision at all.
    without = await reports().create(case["id"],payload)
    assert without["snapshot"]["laboratory"] == []
    skipped = await lab().decide(case["id"],decision)
    assert skipped["status"] == "demo_skipped" and skipped["task_id"] is None
    assert (await lab().decide(case["id"],decision))["reused"]
    assert (await reports().get(without["id"]))["effective_status"] == "historical_snapshot"
    first = await reports().create(case["id"],payload)
    assert first["id"] == (await reports().create(case["id"],payload))["id"]
    assert first["id"] != without["id"]
    assert first["snapshot"]["assessment_id"] == assessment["id"] and "task_id" not in first["snapshot"]
    assert first["snapshot"]["llm_called"] is False
    assert not await db.scalar(select(InspectionTask.id))
    assert (await reports("expert2").get(first["id"]))["snapshot"] == first["snapshot"]


@pytest.mark.asyncio
async def test_validation_report_return_keeps_same_case_and_invalidates_old_review(completion):
    db, service, wf, case, source, assessment, product, category, lab, reports, decision, payload = completion
    link = await lab().decide(case["id"],decision.model_copy(update={"mode":"verify"}))
    old_report = await reports().create(case["id"],payload)
    source2 = await service("expert").create_event(QualityEventCreate(record_type="inspection_report",
        external_record_id="DEMO-VALIDATION",content={"text":"【模拟验证资料】没有真实读数，补充使用条件描述"},
        product_ref={"product_id":product,"category_id":category}, data_nature="synthetic"))
    request = LabReturn(expected_revision=1,record_ids=[source2["id"]],note="明确为模拟验证材料，返回原事项重评")
    returned = await lab().return_records(link["id"],request)
    assert returned["added_new_evidence"] and returned["status"] == "demo_returned"
    assert (await lab().return_records(link["id"],request))["reused"]
    assert returned["assessment_id"] == assessment["id"]
    workspace = await service("expert").risk_workspace(case["id"])
    assert len(workspace["records"]) == 2 and len(workspace["assessment_versions"]) == 1
    assert workspace["latest_assessment"]["effective_status"] == "needs_reassessment"
    historical = await reports().get(old_report["id"])
    assert historical["snapshot"] == old_report["snapshot"] and historical["effective_status"] == "historical_snapshot"
    with pytest.raises(ConflictError): await reports().create(case["id"],payload)
    newer = await submitted(wf(),case,source,base=assessment["id"])
    await wf("expert2").review(case["id"],newer["id"],DemoDraftReview(expected_revision=3,decision="approve",comment="新版独立核对模拟资料"))
    report2 = await reports().create(case["id"],RiskReportCreate(assessment_id=newer["id"],expected_basis_hash=newer["demo_validation"]["basis_hash"]))
    assert report2["snapshot"]["assessment_version"] == 2 and len(report2["snapshot"]["records"]) == 2
    assert report2["effective_status"] == "demo_frozen"
    assert len(await reports().list(case["id"])) == 2
    assert len(list(await db.scalars(select(QualityRiskCase)))) == 1


@pytest.mark.asyncio
async def test_existing_task_link_checks_object_demo_marker_and_never_starts_it(completion):
    db, service, _, case, _, _, product, _, lab, _, decision, _ = completion
    task = InspectionTask(id=str(uuid4()),org_id=service("expert").org_id,created_by=service("expert").actor,
        product_id=product,quality_product_id=product,spec_code="demo-unverified",image_urls=[],status="pending",meta_data={"execution_mode":"demo"})
    db.add(task);await db.flush()
    from app.models.result import InspectionResult
    from app.services.knowledge_source_service import KnowledgeSourceService
    result=InspectionResult(id=str(uuid4()),org_id=service('expert').org_id,task_id=task.id,verdict='uncertain',overall_score=0,llm_model='manual-demo')
    db.add(result);await db.flush()
    with pytest.raises(ValidationError):await KnowledgeSourceService(db,service('expert').current).resolve('task_result',result.id)
    link = await lab().decide(case["id"],decision.model_copy(update={"mode":"verify","task_id":task.id}))
    assert link["task_id"] == task.id and task.status == "pending" and task.started_at is None
    assert len((await lab().options(case["id"]))["tasks"]) == 1
    with pytest.raises(ValidationError): await lab().decide(case["id"],decision.model_copy(update={"task_id":task.id}))
    task.meta_data={};await db.flush()
    assert (await lab().list(case["id"]))[0]["issues"]
    with pytest.raises(ValidationError): await lab().decide(case["id"],decision.model_copy(update={"mode":"verify","task_id":task.id}))


@pytest.mark.asyncio
async def test_unknown_objects_wrong_types_revision_and_frozen_return_are_rejected(completion):
    _, service, _, case, source, _, product, category, lab, _, decision, _ = completion
    skipped = await lab().decide(case["id"],decision)
    with pytest.raises(ConflictError): await lab().return_records(skipped["id"],LabReturn(expected_revision=1,record_ids=[source["id"]],note="尚未选择验证"))
    with pytest.raises(ConflictError): await lab().decide(case["id"],decision.model_copy(update={"mode":"verify","expected_revision":99}))
    link = await lab().decide(case["id"],decision.model_copy(update={"mode":"verify","expected_revision":1}))
    unknown = await service("expert").create_event(QualityEventCreate(record_type="inspection_report",content={"text":"模拟资料，对象未知"}))
    for rid in [unknown["id"], source["id"], str(uuid4())]:
        with pytest.raises((ValidationError,NotFoundError)):
            await lab().return_records(link["id"],LabReturn(expected_revision=2,record_ids=[rid],note="不符合条件"))
    known = await service("expert").create_event(QualityEventCreate(record_type="inspection_report",content={"text":"模拟同对象资料"},product_ref={"product_id":product,"category_id":category}))
    await lab().return_records(link["id"],LabReturn(expected_revision=2,record_ids=[known["id"]],note="正确关联"))
    with pytest.raises(ConflictError): await lab().return_records(link["id"],LabReturn(expected_revision=3,record_ids=[known["id"]],note="不能覆盖固定的回流说明"))


@pytest.mark.asyncio
async def test_report_and_validation_role_scope_withdrawal_and_snapshot_integrity(completion):
    db, service, _, case, source, _, _, _, lab, reports, decision, payload = completion
    first = await reports().create(case["id"],payload)
    for role in ["user","admin"]:
        with pytest.raises(ForbiddenError): await reports(role).get(first["id"])
        with pytest.raises(ForbiddenError): await lab(role).list(case["id"])
    outsider = str(uuid4())
    db.add(User(id=outsider,org_id=service("expert").org_id,username="other-expert",email="other@test.invalid",role="expert",is_active=True));await db.flush()
    other = RiskReportService(db,CurrentUser(user_id=outsider,org_id=service("expert").org_id,role="expert"))
    with pytest.raises(ForbiddenError): await other.get(first["id"])
    with pytest.raises(ForbiddenError): await other.create(case["id"],payload)
    with pytest.raises(NotFoundError): await reports().get(str(uuid4()))
    (await db.get(QualitySourceRecord,source["id"])).deleted_at=datetime.utcnow();await db.flush()
    with pytest.raises(NotFoundError): await reports().get(first["id"])
    assert await reports().list(case["id"]) == []


@pytest.mark.asyncio
@pytest.mark.parametrize("format",["pdf","docx"])
async def test_render_independent_reports_contains_boundary_sources_review_and_unknowns(completion,format):
    _, _, _, case, _, _, _, _, lab, reports, decision, payload = completion
    await lab().decide(case["id"],decision)
    report = await reports().create(case["id"],payload)
    content = build_renderer(format).render({"risk_report":report,"report_name":report["title"],"report_type_label":"独立演示研判报告"})
    if format=='pdf':
        from pypdf import PdfReader
        assert content.startswith(b'%PDF')
        output='\n'.join(p.extract_text() for p in PdfReader(BytesIO(content)).pages)
    else:
        from docx import Document
        doc=Document(BytesIO(content));output='\n'.join(p.text for p in doc.paragraphs)
    for label in ["演示研判报告","模拟资料","独立复核","抽查建议","跳过验证","发生概率：未知"]:
        assert label in output
    assert "PIAP 任务检测记录" not in output


@pytest.mark.asyncio
async def test_export_listing_and_direct_download_cannot_leak_other_readers(completion):
    db, service, _, case, _, _, _, _, _, reports, _, payload = completion
    report=await reports().create(case["id"],payload)
    job=ExportJob(id=str(uuid4()),org_id=service("expert").org_id,actor_id=service("expert").actor,report_name=report["title"],
        report_type="risk_assessment",format="pdf",status="success",config_json=json.dumps({"risk_report_id":report["id"],"snapshot_hash":report["snapshot_hash"]}),
        expires_at=datetime.utcnow()+timedelta(days=7),created_at=datetime.utcnow(),updated_at=datetime.utcnow())
    db.add(job);await db.flush()
    def exports(role):return ExportJobService(db,service(role).org_id,actor_user_id=service(role).actor,actor_role=service(role).current.role)
    assert (await exports("expert").list_jobs(1,20))[0] == 1
    assert (await exports("expert2").list_jobs(1,20))[0] == 0
    with pytest.raises(ForbiddenError): await exports("expert2").get_detail(job.id)
    assert (await exports("expert").get_detail(job.id)).source_version == report["snapshot_hash"]
    with pytest.raises(ValidationError):
        await exports("expert").create_job(service("expert").actor,{"report_type":"risk_assessment","format":"pdf","config_json":{"risk_report_id":report["id"],"task_id":"invented"}})


@pytest.mark.asyncio
async def test_simulated_validation_material_cannot_be_admitted_as_formal_knowledge(completion):
    db,service,_,_,source,_,_,_,_,_,_,_=completion
    from app.services.knowledge_source_service import KnowledgeSourceService
    resolver=KnowledgeSourceService(db,service('expert').current)
    with pytest.raises(ValidationError):await resolver.resolve('intake_record',source['id'])
    row=await db.get(QualitySourceRecord,source['id'])
    row.data_nature='observed';row.provenance={'execution_mode':'demo'};await db.flush()
    with pytest.raises(ValidationError):await resolver.resolve('intake_record',source['id'])
    row.provenance={};row.content={'text':'observed fixture'};await db.flush()
    assert (await resolver.resolve('intake_record',source['id']))['authority']=='observed_material'


@pytest.mark.asyncio
async def test_report_with_current_suggestion_is_current_then_new_feedback_is_historical(sampling):
    db,service,_,case,_,assessment,_,_,suggestions,create=sampling
    from app.schemas.demo_sampling import DemoSamplingSave,DemoSamplingConfirm
    from test_demo_sampling import content
    suggestion=await suggestions().create(case['id'],create)
    await suggestions().mutate(suggestion['id'],DemoSamplingSave(expected_revision=1,content=content()))
    await suggestions().mutate(suggestion['id'],DemoSamplingConfirm(expected_revision=2),confirm=True)
    reports=RiskReportService(db,service('expert').current)
    # The assessment basis and the full collaboration context have different hashes.
    from app.models.quality_risk import QualityRiskAssessment
    row=await db.get(QualityRiskAssessment,assessment['id'])
    report=await reports.create(case['id'],RiskReportCreate(assessment_id=assessment['id'],expected_basis_hash=row.model_versions['demo_workflow']['basis_hash']))
    assert report['effective_status']=='demo_frozen'
    assert (await reports.get(report['id']))['effective_status']=='demo_frozen'
    from app.services.meeting_service import MeetingService
    link=(await suggestions().collab.list_for_case(case['id']))[0]
    await MeetingService(db,service('expert').org_id,service('expert').actor,role='expert').send_message(link['room_id'],'新的公开讨论意见')
    historical=await reports.get(report['id'])
    assert historical['effective_status']=='historical_snapshot' and historical['snapshot']==report['snapshot']
