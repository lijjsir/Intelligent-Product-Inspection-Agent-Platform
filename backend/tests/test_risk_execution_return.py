"""Native execution -> original risk -> new review, using only isolated test fixtures."""
import copy
import hashlib
from io import BytesIO
from datetime import datetime
from uuid import uuid4

import pytest
import pytest_asyncio
from pypdf import PdfWriter
from sqlalchemy import DateTime, DefaultClause, MetaData, func, select, text

from test_quality_risk_v4 import domain  # noqa: F401
from test_risk_sampling_workflow import flow  # noqa: F401
from test_formal_risk_completion import formal, collaboration  # noqa: F401
from app.core.exceptions import ConflictError, ForbiddenError, NotFoundError, ValidationError
from app.models.base import Base
from app.models.inspection_execution import InspectionExecutionRun
from app.models.export_job import ExportJob
from app.models.inspection_standard_library import InspectionStandardLibrary, StandardDocument
from app.models.quality_risk import QualityAttachment, QualityRiskCase, QualitySourceRecord
from app.models.standard_admission import StandardDocumentRevision
from app.models.task import InspectionTask
from app.schemas.quality_risk import ProductCategoryCreate, ProductCreate, QualityEventCreate, RiskAnalyzeRequest, RiskCaseCreate, RiskReviewRequest
from app.schemas.risk_completion import LabDecision, LabExecutionReturn, LabExecutionWithdraw, RiskReportCreate
from app.services.inspection_execution_service import InspectionExecutionService
from app.services.risk_decision_context_service import RiskDecisionContextService
from app.services.risk_lab_service import RiskLabService
from app.services.risk_report_service import RiskReportService
from app.services.risk_execution_evidence_service import execution_returns


@pytest_asyncio.fixture
async def returning(formal, monkeypatch):
    db, risk, supervision, record, _, _ = formal
    metadata = MetaData()
    for name in ('inspection_specs', 'spec_change_logs', 'product_batches', 'export_jobs'):
        table = Base.metadata.tables[name].to_metadata(metadata)
        for column in table.c:
            if isinstance(column.type, DateTime) and column.server_default is not None:
                column.server_default = DefaultClause(text('CURRENT_TIMESTAMP'))
    async with db.bind.begin() as conn:
        await conn.run_sync(metadata.create_all)
    writer = PdfWriter()
    writer.add_blank_page(width=500, height=700)
    output = BytesIO()
    writer.write(output)
    class Storage:
        content = output.getvalue()
        def get_bytes(self, **kwargs): return (self.content, 'application/pdf') if self.content else None
    storage = Storage()
    monkeypatch.setattr('app.services.quality_risk_service.build_object_storage', lambda:storage)
    monkeypatch.setattr('app.services.execution_evidence_service.build_object_storage', lambda:storage)
    category = await risk('admin').create_category(ProductCategoryCreate(code='RETURN-TEST', name='隔离回流测试产品'))
    product = await risk('admin').create_product(ProductCreate(category_id=category['id'], name='隔离测试产品'))
    case = await risk('expert').create_risk_case(RiskCaseCreate(title='隔离检测回流事项', scope_type='product',
        scope={'product_id':product['id'], 'product_category':'隔离回流测试产品'}, source_record_ids=[record['id']]))
    assessment = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离原始资料独立复核', risk_level='high'))
    task = InspectionTask(org_id=risk('expert').org_id, created_by=risk('user').actor, product_id=product['id'],
        quality_product_id=product['id'], product_category_id=category['id'], spec_code='ISOLATED-RETURN', image_urls=[], status='pending',
        meta_data={'input_mode':'report', 'product_category_name':'隔离回流测试产品', 'production_date':'2026-01-01'})
    library = InspectionStandardLibrary(id=str(uuid4()), org_id=task.org_id, name='隔离测试依据，非正式标准', product_family='隔离回流测试产品', rag_space_ids=[])
    attachment = QualityAttachment(org_id=task.org_id, uploaded_by=risk('expert').actor, file_name='isolated-fixture.pdf',
        mime_type='application/pdf', size_bytes=len(storage.content), sha256=hashlib.sha256(storage.content).hexdigest(), bucket='fixture', object_key='fixture')
    db.add_all([task, library, attachment])
    await db.flush()
    report = await risk('expert').create_event(QualityEventCreate(record_type='inspection_report',
        content={'text':'隔离测试原件，仅验证软件流程，不是真实检测报告'}, product_ref={'product_id':product['id']}, attachment_ids=[attachment.id]))
    doc = StandardDocument(org_id=task.org_id, library_id=library.id, standard_no='TEST-QS-RETURN', standard_name='隔离软件测试依据',
        standard_level='TEST', domain='隔离回流测试产品', product_category='隔离回流测试产品', file_name='fixture', file_path='fixture')
    db.add(doc)
    await db.flush()
    candidate = {'id':str(uuid4()), 'candidate_type':'inspection_requirement', 'status':'confirmed',
        'reviewed_by':risk('expert2').actor, 'indicator':'测试指标', 'measurement_key':'test_value', 'unit':'unit', 'method':'fixture',
        'operator':'le', 'numeric_upper':10.0, 'execution_mode':'numeric', 'expected_conditions':{}, 'clause_no':'test', 'page_number':1}
    revision = StandardDocumentRevision(org_id=task.org_id, document_id=doc.id, revision_no=1, processing_status='parsed',
        review_status='approved', publication_status='published', file_hash='isolated-fixture', candidates_json=[candidate], pages_json=[],
        snapshot={'standard_no':'TEST-QS-RETURN', 'product_category':'隔离回流测试产品', 'scope_notes':'仅隔离软件测试',
            'standard_status':'现行', 'production_from':'2025-01-01', 'production_to':'2027-12-31'})
    db.add(revision)
    await db.flush()
    doc.published_revision_id = revision.id
    # Native task association freezes the source, original file and sample mapping.
    from app.services.inspection_execution_context_service import InspectionExecutionContextService
    context_service = InspectionExecutionContextService(db, risk('expert').current, storage)
    context = await context_service.context(task)
    linked = await context_service.link_report(task, {'request_key':str(uuid4()), 'expected_context_hash':context['context_hash'],
        'source_record_id':report['id'], 'attachment_id':attachment.id, 'sample_codes':['fixture-A'], 'reason':'隔离软件测试对象对应核对'})
    sample = linked['reports'][0]['samples'][0]['id']
    payload = {'request_key':str(uuid4()), 'task_id':task.id, 'session_id':None, 'product_category':'隔离回流测试产品',
        'production_date':'2026-01-01', 'policy_id':None, 'requirements':[{'revision_id':revision.id, 'candidate_id':candidate['id']}],
        'sample_ids':[sample], 'observations':[{'kind':'report_attachment', 'id':attachment.id, 'sample_id':sample,
        'measurement_key':'test_value', 'value':11.0, 'unit':'unit', 'method':'fixture', 'conditions':{}, 'report_page':1, 'report_item':'fixture'}]}
    async def native_result():
        body = {**copy.deepcopy(payload), 'request_key':str(uuid4())}
        native = InspectionExecutionService(db, risk('expert').current, storage)
        created = await native.create(body)
        assert not created['current_source_issues'], created
        run = await native.get(created['id'])
        await InspectionExecutionService(db, risk('expert2').current, storage).review(run.id, {'expected_version':run.row_version,
            'request_key':str(uuid4()), 'decision':'accept', 'reason':'隔离逐项复核，仅软件测试', 'report_values_verified':True,
            'evidence_refs':['fixture page1'], 'verified_evidence_ids':[o['evidence_id'] for o in run.snapshot['observations']]})
        return run
    run = await native_result()
    lab = RiskLabService(db, risk('expert').current)
    workspace = await risk('expert').risk_workspace(case['id'])
    link = await lab.decide(case['id'], LabDecision(assessment_id=assessment['id'], expected_basis_hash=workspace['workflow_context']['basis_hash'],
        mode='verify', purpose='隔离验证指标范围', task_id=task.id))
    candidate = (await lab.execution_candidates(link['id']))['items'][0]
    assert candidate['eligible'], candidate
    request = LabExecutionReturn(expected_revision=link['revision'], run_id=run.id, expected_run_hash=candidate['definition_hash'], note='隔离已复核结果返回原事项')
    await db.commit()
    return db, risk, supervision, case, assessment, task, report, run, lab, link, request, native_result


@pytest.mark.asyncio
async def test_native_result_returns_once_and_new_assessment_consumes_it(returning):
    db, risk, _, case, assessment, task, report, run, lab, link, request, _ = returning
    old = copy.deepcopy(await risk('expert').get_assessment(assessment['id']))
    cases_before = await db.scalar(select(func.count()).select_from(QualityRiskCase))
    sources_before = await db.scalar(select(func.count()).select_from(QualitySourceRecord))
    response = await lab.return_execution(link['id'], request)
    duplicate = await lab.return_execution(link['id'], request)
    assert duplicate['reused'] and duplicate['return_id'] == response['return_id']
    with pytest.raises(ConflictError):
        await lab.return_execution(link['id'], request.model_copy(update={'note':'不得静默改写原说明'}))
    assert len(await execution_returns(db, task.org_id, case['id'])) == 1
    assert await db.scalar(select(func.count()).select_from(QualityRiskCase)) == cases_before
    assert await db.scalar(select(func.count()).select_from(QualitySourceRecord)) == sources_before
    stale = await risk('expert').get_assessment(assessment['id'])
    assert stale['effective_status'] == 'needs_reassessment'
    assert stale['review_comment'] == old['review_comment']
    with pytest.raises(ValidationError):
        await RiskDecisionContextService(db, risk('expert').current).accepted(case['id'])
    new = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    assert new['version'] == assessment['version']+1 and new['probability'] is None
    assert new['model_versions']['returned_execution_refs'][0]['run_id'] == run.id
    assert any(m.get('run_id') == run.id for m in new['standard_matches'])
    with pytest.raises(ForbiddenError):
        await risk('expert').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='禁止自审', risk_level='high'))
    reviewed = await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离新版检测范围独立核对', risk_level='high'))
    workspace = await risk('expert').risk_workspace(case['id'])
    result = await RiskReportService(db, risk('expert').current).create(case['id'], RiskReportCreate(
        assessment_id=reviewed['id'], expected_basis_hash=workspace['workflow_context']['basis_hash']))
    assert result['snapshot']['execution_results'][0]['ref']['run_id'] == run.id
    assert '已复核检测结果回流' in str(result['sections'])
    assert result['snapshot']['execution_results'][0]['summary']['product_verdict'] == 'not_assessed'
    from app.services.report_renderers.docx_renderer import DocxRenderer
    from app.services.report_renderers.pdf_renderer import PdfRenderer
    from docx import Document
    from pypdf import PdfReader
    document = {'risk_report':result, 'report_name':result['title']}
    pdf, docx = PdfRenderer().render(document), DocxRenderer().render(document)
    assert run.id in '\n'.join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    text = '\n'.join(p.text for p in Document(BytesIO(docx)).paragraphs)
    assert '已复核检测结果回流' in text and '单任务检测记录' not in text
    import os
    if os.environ.get('PIAP_WRITE_ACCEPTANCE_ARTIFACTS') == '1':
        from pathlib import Path
        folder = Path(__file__).resolve().parents[2]/'.codex_tmp/piap_quality_demo'
        folder.mkdir(parents=True, exist_ok=True)
        (folder/'execution-return-sample.pdf').write_bytes(pdf)
        (folder/'execution-return-sample.docx').write_bytes(docx)


@pytest.mark.asyncio
@pytest.mark.parametrize('change', ['draft', 'self_review', 'other_task', 'other_product', 'private', 'source_removed', 'file_missing', 'source_changed', 'standard_removed', 'stale_hash', 'stale_revision', 'demo_task'])
async def test_return_rejects_invalid_results_without_mutating_case(returning, change):
    db, risk, _, case, _, task, report, run, lab, link, request, _ = returning
    source = await db.get(QualitySourceRecord, report['id'])
    if change == 'draft': run.status = 'draft'
    elif change == 'self_review': run.reviewed_by = run.created_by
    elif change == 'other_task': run.task_id = str(uuid4())
    elif change == 'other_product': task.quality_product_id = str(uuid4())
    elif change == 'private': source.authorization_scope = 'private'
    elif change == 'source_removed': source.deleted_at = datetime.utcnow()
    elif change == 'file_missing': risk('expert').storage.content = None
    elif change == 'source_changed': source.content = {'text':'原报告资料已更新'}
    elif change == 'standard_removed':
        revision = await db.get(StandardDocumentRevision, run.snapshot['source_refs'][0]['revision_id'])
        revision.publication_status = 'unpublished'
    elif change == 'stale_hash': request = request.model_copy(update={'expected_run_hash':'b'*64})
    elif change == 'stale_revision': request = request.model_copy(update={'expected_revision':999})
    elif change == 'demo_task': task.meta_data = {**task.meta_data, 'execution_mode':'demo'}
    await db.flush()
    with pytest.raises((ValidationError, ForbiddenError, NotFoundError, ConflictError)):
        await lab.return_execution(link['id'], request)
    assert not await execution_returns(db, task.org_id, case['id'])
    assert (await risk('expert').risk_case(case['id'])).status == 'risk_assessed'


@pytest.mark.asyncio
async def test_multiple_results_are_not_lost_and_stale_writes_are_rejected(returning):
    db, risk, _, case, _, task, _, run, lab, link, request, native_result = returning
    first = await lab.return_execution(link['id'], request)
    second_run = await native_result()
    choices = await lab.execution_candidates(link['id'], page=1, size=1)
    assert choices['total'] == 2 and len(choices['items']) == 1
    candidate = (await lab.execution_candidates(link['id']))['items'][0]
    assert candidate['id'] == second_run.id
    second_request = LabExecutionReturn(expected_revision=link['revision'], run_id=second_run.id,
        expected_run_hash=candidate['definition_hash'], note='隔离另一执行版本')
    with pytest.raises(ConflictError): await lab.return_execution(link['id'], second_request)
    await lab.return_execution(link['id'], second_request.model_copy(update={'expected_revision':first['revision']}))
    new = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    assert {r['run_id'] for r in new['model_versions']['returned_execution_refs']} == {run.id, second_run.id}


@pytest.mark.asyncio
async def test_withdrawal_is_explicit_and_historical_review_is_never_overwritten(returning):
    db, risk, _, case, old, task, _, run, lab, link, request, _ = returning
    returned = await lab.return_execution(link['id'], request)
    draft = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    accepted = await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离回流新版复核', risk_level='high'))
    workspace = await risk('expert').risk_workspace(case['id'])
    report_service = RiskReportService(db, risk('expert').current)
    report = await report_service.create(case['id'], RiskReportCreate(assessment_id=draft['id'], expected_basis_hash=workspace['workflow_context']['basis_hash']))
    frozen = copy.deepcopy(report['snapshot'])
    run.status = 'awaiting_evidence'
    await db.flush()
    assert (await risk('expert').get_assessment(accepted['id']))['effective_status'] == 'needs_reassessment'
    assert (await report_service.get(report['id']))['effective_status'] == 'historical_snapshot'
    with pytest.raises(ValidationError): await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    with pytest.raises(ForbiddenError):
        await RiskLabService(db, risk('expert2').current).withdraw_execution(link['id'], returned['return_id'],
            LabExecutionWithdraw(expected_revision=returned['revision'], note='非操作人员不能撤回'))
    result = await lab.withdraw_execution(link['id'], returned['return_id'], LabExecutionWithdraw(expected_revision=returned['revision'], note='原结果待补证，撤回本次引用'))
    assert result['execution_returns'][0]['withdrawn_at'] and result['execution_returns'][0]['note'] == request.note
    assert not await execution_returns(db, task.org_id, case['id'])
    assert (await report_service.get(report['id']))['snapshot'] == frozen
    replacement = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    assert replacement['version'] == old['version']+2 and not replacement['model_versions']['returned_execution_refs']
    assert (await risk('expert').get_assessment(accepted['id']))['review_comment'] == accepted['review_comment']


@pytest.mark.asyncio
async def test_revoked_report_access_masks_history_and_blocks_frozen_report(returning):
    db, risk, _, case, _, _, source, _, lab, link, request, _ = returning
    returned = await lab.return_execution(link['id'], request)
    new = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离共享来源复核', risk_level='high'))
    workspace = await risk('expert').risk_workspace(case['id'])
    svc = RiskReportService(db, risk('expert').current)
    report = await svc.create(case['id'], RiskReportCreate(assessment_id=new['id'], expected_basis_hash=workspace['workflow_context']['basis_hash']))
    record = await db.get(QualitySourceRecord, source['id'])
    record.authorization_scope = 'private'
    await db.flush()
    with pytest.raises(ForbiddenError): await svc.get(report['id'])
    listed = (await lab.list(case['id']))[0]
    assert listed['execution_returns'][0]['summary'] is None and request.note not in str(listed)
    assert (await lab.execution_candidates(link['id']))['items'][0]['eligible'] is False
    # Withdrawal remains available when the inaccessible result can no longer be reused.
    await lab.withdraw_execution(link['id'], returned['return_id'], LabExecutionWithdraw(expected_revision=listed['revision'], note='权限撤回后移除引用'))


@pytest.mark.asyncio
async def test_result_cannot_cross_organization_or_role_boundaries(returning):
    db, risk, _, _, _, _, _, run, _, link, request, _ = returning
    for role in ('user', 'admin', 'algorithm_engineer'):
        with pytest.raises(ForbiddenError):
            await RiskLabService(db, risk(role).current).return_execution(link['id'], request)
    current = risk('expert').current.model_copy(update={'org_id':str(uuid4())})
    with pytest.raises((ForbiddenError, NotFoundError)):
        await RiskLabService(db, current).return_execution(link['id'], request)


@pytest.mark.asyncio
async def test_returned_result_reaches_model_input_without_claiming_rule_output_is_a_model(returning, monkeypatch):
    db, risk, _, case, _, _, _, run, lab, link, request, _ = returning
    await lab.return_execution(link['id'], request)
    complaint = await risk('expert').create_event(QualityEventCreate(record_type='consumer_complaint', content={'text':'隔离原始线索'}, occurred_at=None))
    await risk('expert').append_risk_sources(case['id'], [complaint['id']])
    from app.services.supervision_service import SupervisionService
    from agent.subgraphs.supervision.agents import AGENTS
    captured = []
    async def execute(self, snapshot, agent, **kwargs):
        captured.append(copy.deepcopy(snapshot))
        return await AGENTS[agent].run(snapshot)
    async def no_model(self, *args): return None
    monkeypatch.setattr(SupervisionService, 'execute_agent_via_manager', execute)
    monkeypatch.setattr(SupervisionService, 'model_call', no_model)
    result = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    assert captured[0]['data']['execution_results'][0]['run_id'] == run.id
    observed = captured[0]['data']['evidence']
    assert any(e['evidence_id'] == f'execution:{run.id}:{run.row_version}' and e['occurred_at'] is None for e in observed)
    assert result['model_versions']['llm_called'] is False and result['probability'] is None


@pytest.mark.asyncio
async def test_http_return_and_withdraw_commit_and_validate_input(returning):
    _, _, _, _, _, _, _, _, lab, link, request, _ = returning
    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient
    from app.api.v1.quality_risk import router, risk_lab_service
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[risk_lab_service] = lambda:lab
    async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
        path = f"/risk-laboratory-verifications/{link['id']}/execution-results"
        assert (await client.get(path+'?page=0')).status_code == 422
        assert (await client.post(path, json={'run_id':str(uuid4())})).status_code == 422
        response = await client.post(path, json=request.model_dump(mode='json'))
        assert response.status_code == 200, response.text
        returned = response.json()['data']
        assert (await client.get(path)).json()['data']['items'][0]['already_returned']
        result = await client.post(path+f"/{returned['return_id']}/withdraw", json={'expected_revision':returned['revision'], 'note':'隔离接口撤回引用'})
        assert result.status_code == 200, result.text
        assert result.json()['data']['status'] == 'reference_withdrawn'


@pytest.mark.asyncio
async def test_new_version_updates_suggestions_and_report_without_rewriting_old_snapshots(returning):
    db, risk, supervision, case, old, task, report_source, run, lab, link, request, _ = returning
    from app.services.demo_sampling_service import DemoSamplingService
    from app.schemas.demo_sampling import DemoSamplingCreate, DemoSamplingSave, DemoSamplingConfirm, DemoSamplingContent
    suggestions = DemoSamplingService(db, risk('expert').current)
    reports = RiskReportService(db, risk('expert').current)
    async def completed_suggestion(assessment):
        await collaboration((db, risk, supervision, report_source, case, assessment))
        candidate = (await suggestions.candidates(case['id']))['items'][0]
        row = await suggestions.create(case['id'], DemoSamplingCreate(assessment_id=assessment['id'],
            collaboration_id=candidate['snapshot']['collaboration_id'], expected_context_hash=candidate['context_hash']))
        row = await suggestions.mutate(row['id'], DemoSamplingSave(expected_revision=row['revision'],
            content=DemoSamplingContent(rationale='隔离证据核对', focus='核查项目范围', limitations='概率未校准，执行安排待确认')))
        return await suggestions.mutate(row['id'], DemoSamplingConfirm(expected_revision=row['revision']), confirm=True)
    async def frozen_report(assessment):
        workspace = await risk('expert').risk_workspace(case['id'])
        return await reports.create(case['id'], RiskReportCreate(assessment_id=assessment['id'], expected_basis_hash=workspace['workflow_context']['basis_hash']))
    old_suggestion = await completed_suggestion(old)
    old_report = await frozen_report(old)
    frozen = copy.deepcopy(old_report['snapshot'])
    await lab.return_execution(link['id'], request)
    assert not (await suggestions.get(old_suggestion['id']))['is_current']
    assert (await reports.get(old_report['id']))['effective_status'] == 'historical_snapshot'
    new = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离新版独立确认', risk_level='high'))
    new_suggestion = await completed_suggestion(new)
    assert new_suggestion['snapshot']['assessment_id'] == new['id'] and new_suggestion['is_current']
    # A second validation link to the same result must reuse the original return.
    workspace = await risk('expert').risk_workspace(case['id'])
    second_link = await lab.decide(case['id'], LabDecision(assessment_id=new['id'], expected_basis_hash=workspace['workflow_context']['basis_hash'],
        mode='verify', purpose='同一结果不得重复回流', task_id=task.id))
    repeated = await lab.return_execution(second_link['id'], request)
    assert repeated['id'] == link['id'] and repeated['reused']
    assert len(await execution_returns(db, task.org_id, case['id'])) == 1
    current = await frozen_report(new)
    assert current['effective_status'] == 'frozen' and current['id'] != old_report['id']
    assert any(s['id'] == new_suggestion['id'] for s in current['snapshot']['sampling'])
    assert (await reports.get(old_report['id']))['snapshot'] == frozen


@pytest.mark.asyncio
async def test_formal_export_uses_correct_label_and_does_not_reuse_old_demo_labelled_files(returning, monkeypatch):
    import json
    from datetime import timedelta
    from unittest.mock import AsyncMock
    from app.services.export_job_service import ExportJobService
    from app.services.report_generation_service import _type_label
    db, risk, _, case, _, _, _, _, lab, link, request, _ = returning
    await lab.return_execution(link['id'], request)
    new = await risk('expert').analyze_risk(case['id'], RiskAnalyzeRequest())
    await risk('expert2').review_risk(case['id'], RiskReviewRequest(decision='accept', comment='隔离导出来源复核', risk_level='high'))
    workspace = await risk('expert').risk_workspace(case['id'])
    report = await RiskReportService(db, risk('expert').current).create(case['id'], RiskReportCreate(assessment_id=new['id'], expected_basis_hash=workspace['workflow_context']['basis_hash']))
    legacy_id = str(uuid4())
    db.add(ExportJob(id=legacy_id, org_id=risk('expert').org_id, actor_id=risk('expert').actor,
        report_name=report['title'], report_type='risk_assessment', format='pdf', status='success',
        config_json=json.dumps({'risk_report_id':report['id'], 'snapshot_hash':report['snapshot_hash'], 'renderer_version':'risk-report-v1'}),
        created_at=datetime.utcnow(), updated_at=datetime.utcnow(), expires_at=datetime.utcnow()+timedelta(days=1)))
    await db.flush()
    generate = AsyncMock()
    monkeypatch.setattr('app.services.report_generation_service.generate_report', generate)
    exports = ExportJobService(db, risk('expert').org_id, actor_user_id=risk('expert').actor, actor_role='expert')
    payload = {'report_type':'risk_assessment', 'format':'pdf', 'config_json':{'risk_report_id':report['id']}}
    created = await exports.create_job(risk('expert').actor, payload)
    created_id = created.id
    assert created_id != legacy_id and json.loads(created.config_json)['renderer_version'] == 'risk-report-v2'
    assert created.result_scope['label'] == '风险研判报告'
    assert (await exports.create_job(risk('expert').actor, payload)).id == created_id
    generate.assert_awaited_once()
    assert _type_label('risk_assessment', 'business') == '独立风险研判报告'
    assert _type_label('risk_assessment', 'demo') == '独立演示研判报告'
