import json
from contextlib import asynccontextmanager
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from pypdf import PdfReader

from app.services import report_generation_service as reports


@pytest.mark.asyncio
@pytest.mark.parametrize("format", ["pdf", "docx"])
async def test_task_report_generates_real_files_with_effective_verdict_and_citations(monkeypatch, format):
    job = SimpleNamespace(
        id="job-1",org_id="org-1",actor_id='user-1',status='pending',config_json=json.dumps({"task_id": "task-1"}),
        report_name="产品 <A&B> 检测报告", report_type="single_task", format=format,
    )
    result = SimpleNamespace(
        verdict="pass", reviewed_by=None, overall_score=0.95,
        prompt_version="quality_analysis_v1", reasoning_chain={"standard_evaluation": {"summary": "需核对原始证据"}},
        citations=[{"source": "fixture-source", "quote": "fixture-evidence <A&B>"}],
        defects=[], llm_model="test", tokens_used=10, latency_ms=5,
    )
    database = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock(),expire_all=lambda:None)

    @asynccontextmanager
    async def session():
        yield database

    update = AsyncMock()
    monkeypatch.setattr(reports, "get_session", session)
    bundle={'task':{'id':'task-1','product_id':'sample'},'scope':{'stage':'auxiliary_record','label':'需人工复核','boundary':'产品整体未评定'},
        'selection':{'legacy_only':True,'execution_run_id':None},'auxiliary':{'verdict':'manual_required','score_status':'not_calibrated',
        'citations':result.citations,'summary':'需核对原始证据'},'source_refs':[],'risks':[]}
    monkeypatch.setattr(reports,'prepare_report',AsyncMock(return_value=bundle))
    monkeypatch.setattr(reports,'report_actor',AsyncMock(return_value=SimpleNamespace(org_id='org-1')))
    monkeypatch.setattr('app.services.task_result_bundle_service.TaskResultBundleService.recheck',AsyncMock(return_value=bundle))
    monkeypatch.setattr("app.repositories.export_job_repo.ExportJobRepository.get_by_id", AsyncMock(return_value=job))
    monkeypatch.setattr("app.repositories.export_job_repo.ExportJobRepository.update_status", update)
    monkeypatch.setattr("app.repositories.task_repo.TaskRepository.get", AsyncMock(return_value=SimpleNamespace(id="task-1", product_id="sample", spec_code="fixture-standard")))
    monkeypatch.setattr("app.repositories.result_repo.ResultRepository.get_by_task", AsyncMock(return_value=result))
    monkeypatch.setattr("app.repositories.stability_repo.StabilityRepository.get_by_task", AsyncMock(return_value=None))
    saved = []
    storage = SimpleNamespace(
        ensure_bucket=lambda _bucket: None,
        put_bytes=lambda **kwargs: saved.append(kwargs),
    )
    monkeypatch.setattr(reports, "build_object_storage", lambda: storage)

    await reports.generate_report(job_id="job-1", org_id="org-1")

    assert update.call_args.args == ("job-1", "success")
    payload = saved[0]
    assert payload["object_key"].endswith(f"/report.{format}")
    if format == "pdf":
        assert payload["data"].startswith(b"%PDF")
        text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(payload["data"])).pages)
    else:
        document = Document(BytesIO(payload["data"]))
        text = "\n".join([
            *(paragraph.text for paragraph in document.paragraphs),
            *(cell.text for table in document.tables for row in table.rows for cell in row.cells),
        ])
    assert "需人工复核" in text
    assert "未校准" in text
    assert "fixture-source" in text
    assert "fixture-evidence <A&B>" in text
    database.rollback.assert_not_awaited()
