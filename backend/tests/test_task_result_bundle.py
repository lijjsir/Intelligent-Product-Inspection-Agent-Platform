import copy,json,uuid
from io import BytesIO
from contextlib import asynccontextmanager
import pytest,pytest_asyncio
from sqlalchemy import MetaData,DateTime,DefaultClause,text,select
from sqlalchemy.ext.asyncio import async_sessionmaker
from docx import Document
from pypdf import PdfReader
import test_inspection_execution_service as base
from app.models.export_job import ExportJob
from app.models.user import User
from app.models.organization import Organization
from app.models.permission import RolePermission,UserPermission
from app.models.result import InspectionResult
from app.models.quality_risk import QualityRiskAssessment,QualityRiskCase,QualityEvidenceItem,RiskPolicy
from app.models.inspection_execution import InspectionExecutionRun
from app.core.datetime import utcnow
from app.core.exceptions import NotFoundError,ForbiddenError,ConflictError
from app.services.task_result_bundle_service import TaskResultBundleService
from app.services.export_job_service import ExportJobService
from app.services.inspection_execution_service import InspectionExecutionService
from app.services import report_generation_service as generation


@pytest_asyncio.fixture
async def data(monkeypatch):
    generator=base.data.__wrapped__(monkeypatch);data=await anext(generator);db=data[0]
    await db.commit();meta=MetaData()
    models=[ExportJob,User,Organization,RolePermission,UserPermission,InspectionResult,QualityRiskAssessment,QualityRiskCase,QualityEvidenceItem,RiskPolicy]
    for model in models:
        table=model.__table__.to_metadata(meta)
        for col in table.c:
            if isinstance(col.type,DateTime) and col.server_default is not None:col.server_default=DefaultClause(text('CURRENT_TIMESTAMP'))
    async with db.bind.begin() as conn:await conn.run_sync(meta.create_all)
    now=utcnow();db.add(Organization(id=base.ORG,name='isolated',slug=base.ORG,settings={},is_active=True,created_at=now,updated_at=now))
    for actor,role in [(base.AUTHOR,'user'),(base.EXPERT,'expert')]:
        db.add(User(id=actor,org_id=base.ORG,username=role,email=f'{role}@example.invalid',role=role,is_active=True,created_at=now,updated_at=now))
    await db.commit()
    monkeypatch.setattr('worker.tasks.report_generate_task.generate_report_task.delay',lambda payload:None)
    try:yield data
    finally:await generator.aclose()


async def run(data,*,conditions=False,accept=False):
    if conditions:data[2].candidates_json=[{**data[3],'candidate_type':'method_condition'}]
    result=await InspectionExecutionService(data[0],base.current()).create(base.payload(data))
    row=await data[0].get(InspectionExecutionRun,result['id'])
    if accept:await InspectionExecutionService(data[0],base.current(base.EXPERT,'expert')).review(row.id,base.review(row))
    await data[0].commit();return row


def export_service(data):return ExportJobService(data[0],base.ORG,actor_user_id=base.AUTHOR,actor_role='user')


async def job(data,**config):
    return await export_service(data).create_job(base.AUTHOR,{'report_name':'R3 隔离检测记录','report_type':'single_task','format':config.pop('format','pdf'),
        'config_json':{'task_id':data[1].id,**config}})


def legacy(data,**fields):
    row=InspectionResult(id=str(uuid.uuid4()),org_id=base.ORG,task_id=data[1].id,verdict='pass',overall_score=.95,
        reasoning_chain={'score_status':'not_calibrated'},defects=[],citations={'items':[{'source':'fixture-source','quote':'fixture evidence <A&B>'}]},
        llm_model='fixture',prompt_version='quality_analysis_v1',**fields)
    data[0].add(row);return row


@pytest.mark.asyncio
async def test_execution_only_task_has_shared_summary_and_export_option(data):
    row=await run(data,accept=True)
    service=TaskResultBundleService(data[0],base.current())
    result=await service.summary(data[1].id)
    assert result['execution']['id']==row.id and result['auxiliary'] is None and result['can_export']
    assert result['scope']['stage']=='reviewed_selected_items' and result['scope']['product_verdict']=='not_assessed'
    assert (await service.report_tasks())[0]['has_execution_result']
    with pytest.raises(NotFoundError):await TaskResultBundleService(data[0],base.current(org=base.OTHER)).summary(data[1].id)
    with pytest.raises(NotFoundError):await TaskResultBundleService(data[0],base.current(user=base.EXPERT)).summary(data[1].id)


@pytest.mark.asyncio
async def test_condition_and_image_sources_remain_separate(data):
    row=await run(data,conditions=True);old=legacy(data);await data[0].commit()
    result=await TaskResultBundleService(data[0],base.current()).summary(data[1].id)
    assert result['scope']['label']=='试验条件不符合要求'
    assert result['auxiliary']['kind']=='visual_auxiliary' and result['auxiliary']['product_verdict']=='not_assessed'
    assert {s['kind'] for s in result['source_refs']}=={'inspection_execution','task_result'}
    selected=await TaskResultBundleService(data[0],base.current()).summary(data[1].id,legacy_only=True)
    assert selected['execution'] is None and selected['scope']['stage']=='auxiliary_record'


@pytest.mark.asyncio
async def test_export_freezes_server_source_not_client_snapshot_or_latest_row(data):
    row=await run(data)
    service=TaskResultBundleService(data[0],base.current());summary=await service.summary(data[1].id)
    exported=await job(data,report_snapshot={'forged':'product passed'},execution_run_id=row.id,expected_result_token=summary['result_token'])
    snapshot=json.loads(exported.config_json)['report_snapshot']
    assert snapshot['execution']['id']==row.id and snapshot['scope']['product_verdict']=='not_assessed'
    other=await run(data)
    assert (await service.recheck(snapshot))['execution']['id']==row.id
    assert other.id!=row.id


@pytest.mark.asyncio
async def test_stale_result_token_is_rejected_before_job_creation(data):
    row=await run(data);summary=await TaskResultBundleService(data[0],base.current()).summary(data[1].id)
    await InspectionExecutionService(data[0],base.current(base.EXPERT,'expert')).review(row.id,base.review(row))
    with pytest.raises(ConflictError):await job(data,execution_run_id=row.id,expected_result_token=summary['result_token'])
    assert not (await data[0].scalars(select(ExportJob))).all()


@pytest.mark.asyncio
async def test_review_evidence_change_invalidates_report_and_knowledge_version(data):
    from app.models.inspection_execution import InspectionExecutionAction
    from app.services.knowledge_source_service import KnowledgeSourceService
    row=await run(data,accept=True)
    service=TaskResultBundleService(data[0],base.current())
    snapshot=await service.freeze(data[1].id,execution_run_id=row.id)
    before=await KnowledgeSourceService(data[0],base.current()).resolve('inspection_execution',row.id,preview=True)
    action=await data[0].scalar(select(InspectionExecutionAction).where(InspectionExecutionAction.run_id==row.id,
        InspectionExecutionAction.action=='accept'))
    action.evidence_refs=[{'source':'changed-review-provenance'}]
    await data[0].commit()
    with pytest.raises(ConflictError):await service.recheck(snapshot)
    after=await KnowledgeSourceService(data[0],base.current()).resolve('inspection_execution',row.id,preview=True)
    assert before['source_version']!=after['source_version']


@pytest.mark.asyncio
@pytest.mark.parametrize('format',['pdf','docx'])
async def test_real_worker_renders_frozen_condition_scope_and_review_provenance(data,monkeypatch,format):
    row=await run(data,conditions=True,accept=True);exported=await job(data,format=format,execution_run_id=row.id)
    factory=async_sessionmaker(data[0].bind,expire_on_commit=False)
    @asynccontextmanager
    async def sessions():
        async with factory() as db:yield db
    monkeypatch.setattr(generation,'get_session',sessions)
    saved=[]
    class Storage:
        def ensure_bucket(self,bucket):pass
        def put_bytes(self,**payload):saved.append(payload)
    monkeypatch.setattr(generation,'build_object_storage',lambda:Storage())
    await generation.generate_report(job_id=exported.id,org_id=base.ORG)
    assert saved and saved[0]['object_key'].endswith('report.'+format)
    raw=saved[0]['data']
    if format=='pdf':text_content='\n'.join(p.extract_text() for p in PdfReader(BytesIO(raw)).pages)
    else:
        document=Document(BytesIO(raw));text_content='\n'.join([*(p.text for p in document.paragraphs),*(c.text for t in document.tables for r in t.rows for c in r.cells)])
    for fragment in ['产品整体未评定','试验方法与条件核对',row.id,base.EXPERT,'fixture-report page3']:
        assert ''.join(fragment.split()) in ''.join(text_content.split())
    assert '产品合格' not in text_content and '产品不合格' not in text_content
    await data[0].refresh(exported);assert exported.status=='success'


@pytest.mark.asyncio
async def test_withdrawal_before_worker_fails_without_switching_sources(data,monkeypatch):
    row=await run(data,accept=True);exported=await job(data,execution_run_id=row.id)
    data[2].publication_status='withdrawn';await data[0].commit()
    factory=async_sessionmaker(data[0].bind,expire_on_commit=False)
    @asynccontextmanager
    async def sessions():
        async with factory() as db:yield db
    monkeypatch.setattr(generation,'get_session',sessions)
    with pytest.raises(ConflictError):await generation.generate_report(job_id=exported.id,org_id=base.ORG)
    await data[0].refresh(exported);assert exported.status=='failed' and not exported.file_url


@pytest.mark.asyncio
async def test_worker_checks_tenant_before_mutating_job_and_rechecks_actor_permission(data,monkeypatch):
    await run(data);exported=await job(data)
    factory=async_sessionmaker(data[0].bind,expire_on_commit=False)
    @asynccontextmanager
    async def sessions():
        async with factory() as db:yield db
    monkeypatch.setattr(generation,'get_session',sessions)
    with pytest.raises(NotFoundError):await generation.generate_report(job_id=exported.id,org_id=base.OTHER)
    await data[0].refresh(exported);assert exported.status=='pending'
    user=await data[0].get(User,base.AUTHOR);user.is_active=False;await data[0].commit()
    with pytest.raises(ForbiddenError):await generation.generate_report(job_id=exported.id,org_id=base.ORG)
    await data[0].refresh(exported);assert exported.status=='failed'


@pytest.mark.asyncio
async def test_result_list_includes_execution_only_and_knowledge_tracks_review_version(data):
    from app.services.knowledge_source_service import KnowledgeSourceService
    row=await run(data)
    service=TaskResultBundleService(data[0],base.current())
    page=await service.list_results()
    assert page['total']==1 and page['items'][0]['execution_run_id']==row.id
    before=await KnowledgeSourceService(data[0],base.current()).resolve('inspection_execution',row.id,preview=True)
    await InspectionExecutionService(data[0],base.current(base.EXPERT,'expert')).review(row.id,base.review(row))
    after=await KnowledgeSourceService(data[0],base.current()).resolve('inspection_execution',row.id,preview=True)
    assert before['source_version']!=after['source_version'] and after['authority']=='reviewed_selected_items'
    assert after['metadata']['row_version']==row.row_version and after['metadata']['product_verdict']=='not_assessed'


@pytest.mark.asyncio
async def test_risk_summary_only_links_same_execution_and_preserves_unknown_probability(data):
    from app.models.quality_risk import QualitySourceRecord
    from app.services.risk_evidence_validation_service import record_definition,evidence_definition
    row=await run(data,accept=True);db=data[0]
    source_id=data[1].meta_data['execution_inputs']['report_links'][0]['source_record_id'];source=await db.get(QualitySourceRecord,source_id)
    evidence=QualityEvidenceItem(org_id=base.ORG,evidence_code='isolated-risk',source_record_id=source_id,evidence_type='record',text='真实抽检线索',
        structured_data={},nature='observed',locator={},content_hash='fixture')
    db.add(evidence);await db.flush()
    case=QualityRiskCase(org_id=base.ORG,code='RISK-ISOLATED',title='isolated',scope_type='category',scope={'name':'电动自行车'},
        evidence_ids=[evidence.id],source_record_ids=[source_id],created_by=base.AUTHOR)
    db.add(case);await db.flush()
    assessment=QualityRiskAssessment(org_id=base.ORG,risk_case_id=case.id,version=1,scope=case.scope,evidence_ids=[evidence.id],risk_level='high',status='accepted',
        submitted_by=base.AUTHOR,reviewed_by=base.EXPERT,standard_matches=[{'kind':'inspection_execution','run_id':row.id,'row_version':row.row_version}],
        recommendations=['继续抽查'],model_versions={'evidence_basis':{'evidence_ids':[evidence.id],'records':{source.id:record_definition(source)},'evidence':{evidence.id:evidence_definition(evidence)}}})
    db.add(assessment);await db.commit()
    bundle=await TaskResultBundleService(db,base.current(base.EXPERT,'expert')).summary(data[1].id)
    assert len(bundle['risks'])==1 and bundle['risks'][0]['probability'] is None and bundle['risks'][0]['effective_status']=='accepted'
    assert bundle['scope']['product_verdict']=='not_assessed'
    assert not (await TaskResultBundleService(db,base.current()).summary(data[1].id))['risks']


@pytest.mark.asyncio
async def test_source_withdrawal_after_generation_keeps_historical_job_but_marks_current_state(data):
    row=await run(data,accept=True);exported=await job(data,execution_run_id=row.id)
    data[2].publication_status='withdrawn';await data[0].commit()
    details=await export_service(data).get_detail(exported.id)
    assert details.source_issues and details.effective_status=='historical_snapshot'
    assert json.loads(details.config_json)['report_snapshot']['execution']['status']=='accepted'


@pytest.mark.asyncio
async def test_change_during_render_is_rechecked_before_file_upload(data,monkeypatch):
    row=await run(data,accept=True);exported=await job(data,execution_run_id=row.id)
    factory=async_sessionmaker(data[0].bind,expire_on_commit=False)
    @asynccontextmanager
    async def sessions():
        async with factory() as db:yield db
    monkeypatch.setattr(generation,'get_session',sessions)
    original=TaskResultBundleService.recheck;calls=0
    async def recheck(service,bundle):
        nonlocal calls
        calls+=1
        if calls==2:
            data[2].publication_status='withdrawn';await data[0].commit()
        return await original(service,bundle)
    monkeypatch.setattr(TaskResultBundleService,'recheck',recheck)
    saved=[]
    class Storage:
        def ensure_bucket(self,bucket):pass
        def put_bytes(self,**kw):saved.append(kw)
    monkeypatch.setattr(generation,'build_object_storage',lambda:Storage())
    with pytest.raises(ConflictError):await generation.generate_report(job_id=exported.id,org_id=base.ORG)
    assert not saved
