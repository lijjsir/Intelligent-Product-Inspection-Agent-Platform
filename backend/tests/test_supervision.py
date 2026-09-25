from __future__ import annotations

import asyncio
import io
from datetime import datetime
from uuid import uuid4

import pytest
import pytest_asyncio
from sqlalchemy import MetaData, select, DateTime, DefaultClause, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.exceptions import ConflictError, ForbiddenError, ValidationError
from app.models import (
    Base,
    Organization,
    User,
    ProductLine,
    ProductSku,
    ProductBatch,
    InspectionTask,
)
from app.models.supervision import (
    BusinessReview,
    MeasurementRecord,
    SupervisionEvent,
    SupervisionRecord,
    SupervisionRevision,
)
from app.schemas.user import CurrentUser
from app.schemas.supervision import (
    MeasurementIngest,
    RecordCreate,
    RecordUpdate,
    ReviewDecision,
    ReviewRequest,
    RunRequest,
)
from app.services.supervision_service import SupervisionService
from agent.subgraphs.supervision.agents import AGENTS


def uid():
    return str(uuid4())


@pytest_asyncio.fixture
async def domain():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    metadata = MetaData()
    names = [
        n
        for n in Base.metadata.tables
        if n
        in {
            "organizations",
            "users",
            "product_lines",
            "product_skus",
            "product_batches",
            "inspection_tasks",
            "inspection_results",
            "stability_reports",
            "inspection_standard_libraries",
            "standard_documents",
            "standard_document_chunks",
            "inspection_specs",
            "inspection_goals",
            "inspection_evidence_states",
            "inspection_decisions",
            "device_commands",
            "datasets",
            "dataset_samples",
            "audit_logs",
        }
        or n.startswith(("supervision_", "business_", "measurement_", "device_connections"))
    ]
    for name in names:
        table = Base.metadata.tables[name].to_metadata(metadata)
        # MySQL-specific timestamp defaults are irrelevant to domain tests; values are supplied explicitly.
        for column in table.c:
            column.server_default = (
                DefaultClause(text("CURRENT_TIMESTAMP"))
                if isinstance(column.type, DateTime) and column.server_default is not None
                else None
            )
    async with engine.begin() as c:
        await c.run_sync(metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as db:
        org, other_org = uid(), uid()
        actors = {
            role: uid()
            for role in (
                "admin",
                "user",
                "expert",
                "platform_operator",
                "algorithm_engineer",
                "app_developer",
            )
        }
        actors["expert2"] = uid()
        now = datetime.now()
        for oid in (org, other_org):
            db.add(
                Organization(
                    id=oid,
                    name="试点组织",
                    slug=oid,
                    settings={"quality_supervision_enabled": True},
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                )
            )
        for role, actor in actors.items():
            db.add(
                User(
                    id=actor,
                    org_id=org,
                    username=role,
                    email=f"{role}@example.com",
                    role="expert" if role == "expert2" else role,
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                )
            )
        line, sku, batch, task = uid(), uid(), uid(), uid()
        db.add(
            ProductLine(
                id=line,
                org_id=org,
                code="EV",
                name="电动自行车",
                is_active=True,
                created_at=now,
                updated_at=now,
            )
        )
        db.add(
            ProductSku(
                id=sku,
                org_id=org,
                product_line_id=line,
                code="EV-001",
                name="试点车",
                is_active=True,
                created_at=now,
                updated_at=now,
            )
        )
        db.add(
            ProductBatch(
                id=batch,
                org_id=org,
                product_sku_id=sku,
                batch_no="B-001",
                name="试点批次",
                is_active=True,
                created_at=now,
                updated_at=now,
            )
        )
        db.add(
            InspectionTask(
                id=task,
                org_id=org,
                created_by=actors["user"],
                product_sku_id=sku,
                batch_id=batch,
                product_id="EV-001",
                spec_code="EV-TEST",
                image_urls=[],
                status="collecting",
                priority=5,
                meta_data={"input_mode": "measurement"},
                created_at=now,
                updated_at=now,
            )
        )
        await db.commit()

        def service(role="user", oid=None):
            return SupervisionService(
                db,
                CurrentUser(
                    user_id=actors[role],
                    org_id=oid or org,
                    role="expert" if role == "expert2" else role,
                ),
            )

        yield (
            db,
            service,
            {
                "org": org,
                "other_org": other_org,
                "actors": actors,
                "sku": sku,
                "batch": batch,
                "task": task,
            },
        )
    await engine.dispose()


async def create(svc, kind, code, data):
    return await svc.create(kind, RecordCreate(code=code, name=code, data=data))


async def session_setup(domain):
    db, svc, ids = domain
    device = await create(
        svc("admin"),
        "devices",
        "设备-001",
        {
            "device_type": "温升检测",
            "capabilities": ["温升"],
            "calibration_status": "valid",
            "capacity": 10,
        },
    )
    sample = await create(
        svc(),
        "samples",
        "样品-001",
        {
            "task_id": ids["task"],
            "product_sku_id": ids["sku"],
            "batch_id": ids["batch"],
            "sampled_at": "2026-09-15T10:00:00",
        },
    )
    session = await create(
        svc(),
        "inspection-sessions",
        "检测-001",
        {
            "task_id": ids["task"],
            "sample_ids": [sample["id"]],
            "device_ids": [device["id"]],
            "test_items": [
                {
                    "item": "温升",
                    "unit": "C",
                    "method": "试点方法",
                    "upper_limit": 45,
                    "standard_ref": "试点条款1",
                }
            ],
        },
    )
    measurement = {
        "event_id": "event-1",
        "sample_id": sample["id"],
        "device_id": device["id"],
        "item": "温升",
        "measured_at": "2026-09-15T10:10:00",
        "value": 55,
        "unit": "C",
        "method": "试点方法",
        "quality_flag": "valid",
    }
    return device, sample, session, measurement


@pytest.mark.asyncio
async def test_archive_tenant_and_role_access(domain):
    db, svc, ids = domain
    region = await create(svc("admin"), "regions", "CQ", {})
    with pytest.raises(ForbiddenError):
        await create(svc("user"), "regions", "SC", {})
    with pytest.raises(Exception) as e:
        await svc("user", ids["other_org"]).get("regions", region["id"])
    assert "不存在" in str(e.value)
    for role in ("algorithm_engineer", "app_developer"):
        with pytest.raises(ForbiddenError):
            await svc(role).list("risk-cases")


@pytest.mark.asyncio
async def test_case_ownership_and_immutable_revision(domain):
    db, svc, ids = domain
    case = await create(
        svc(),
        "risk-cases",
        "线索-001",
        {"description": "续航下降", "occurred_at": "2026-09-15T10:00:00"},
    )
    other = SupervisionService(
        db, CurrentUser(user_id=ids["actors"]["expert2"], org_id=ids["org"], role="user")
    )
    with pytest.raises(ForbiddenError):
        await other.update("risk-cases", case["id"], RecordUpdate(version=1, name="改名"))
    await svc().update("risk-cases", case["id"], RecordUpdate(version=1, name="新名称"))
    revision = await db.scalar(
        select(SupervisionRevision).where(
            SupervisionRevision.record_id == case["id"], SupervisionRevision.version == 1
        )
    )
    assert revision.snapshot["name"] == "线索-001"
    with pytest.raises(ConflictError):
        await svc().update("risk-cases", case["id"], RecordUpdate(version=1, name="旧版本"))


@pytest.mark.asyncio
async def test_region_cycles_and_distinct_locations(domain):
    db, svc, ids = domain
    cq = await create(svc("admin"), "regions", "重庆", {})
    district = await create(svc("admin"), "regions", "渝北", {"parent_id": cq["id"]})
    with pytest.raises(ValidationError):
        await svc("admin").update(
            "regions", cq["id"], RecordUpdate(version=1, data={"parent_id": district["id"]})
        )
    case = await create(
        svc(),
        "risk-cases",
        "异地线索",
        {
            "description": "续航下降",
            "occurred_at": "2026-09-15T10:00:00",
            "production_region_id": cq["id"],
            "complaint_region_id": district["id"],
        },
    )
    assert case["data"]["production_region_id"] != case["data"]["complaint_region_id"]


@pytest.mark.asyncio
async def test_expert_review_and_self_review_rejection(domain):
    db, svc, ids = domain
    assert await svc("expert").permissions() == []
    record = SupervisionRecord(
        org_id=ids["org"],
        kind="sampling-plans",
        code="P",
        name="P",
        created_by=ids["actors"]["expert"],
        data={"analysis": {"status": "completed", "selected": []}},
        status="awaiting_review",
    )
    db.add(record)
    await db.flush()
    r = BusinessReview(
        org_id=ids["org"],
        record_id=record.id,
        version=1,
        operation="sampling.approve",
        requester_id=ids["actors"]["expert"],
    )
    db.add(r)
    await db.flush()
    with pytest.raises(ForbiddenError):
        await svc("expert").review_decision(r.id, ReviewDecision(decision="accept", comment="同意"))
    await svc("expert2").review_decision(
        r.id, ReviewDecision(decision="accept", comment="独立专家确认")
    )


@pytest.mark.asyncio
async def test_measurement_duplicate_conflict_and_invalid_unit(domain):
    db, svc, ids = domain
    device, sample, session, m = await session_setup(domain)
    req = MeasurementIngest(session_id=session["id"], source_key="batch-1", measurements=[m])
    assert (await svc().ingest(req))["accepted"] == 1
    assert (await svc().ingest(req))["accepted"] == 0
    assert len(list(await db.scalars(select(MeasurementRecord)))) == 1
    altered = {**m, "value": 40}
    with pytest.raises(ConflictError):
        await svc().ingest(
            MeasurementIngest(
                session_id=session["id"], source_key="batch-2", measurements=[altered]
            )
        )
    with pytest.raises(ValidationError):
        await svc().ingest(
            MeasurementIngest(
                session_id=session["id"],
                source_key="batch-3",
                measurements=[{**m, "event_id": "2", "unit": "F"}],
            )
        )
    with pytest.raises(ForbiddenError):
        await svc().ingest(req, device_id=uid())


@pytest.mark.asyncio
async def test_file_preview_validation_and_confirm_version(domain):
    db, svc, ids = domain
    device, sample, session, m = await session_setup(domain)
    headers = ",".join(m.keys())
    row = ",".join(str(v) for v in m.values())
    preview = await svc().preview_file(
        session["id"], "测量.csv", (headers + "\n" + row + "\n").encode(), {}
    )
    assert preview["can_confirm"] and preview["valid_count"] == 1
    result = await svc().confirm_file(session["id"], preview["preview_id"], "file-1")
    assert result["accepted"] == 1
    with pytest.raises(ConflictError):
        await svc().confirm_file(session["id"], preview["preview_id"], "file-2")
    broken = await svc().preview_file(
        session["id"], "bad.csv", "item,value\n温升,bad\n".encode(), {}
    )
    assert not broken["can_confirm"]


@pytest.mark.asyncio
async def test_xlsx_ingest_uses_same_unit_validation(domain):
    from openpyxl import Workbook

    db, svc, ids = domain
    d, s, session, m = await session_setup(domain)
    wb = Workbook()
    wb.active.append(list(m.keys()))
    wb.active.append(list(m.values()))
    buffer = io.BytesIO()
    wb.save(buffer)
    preview = await svc().preview_file(session["id"], "测量.xlsx", buffer.getvalue(), {})
    assert preview["valid_count"] == 1
    assert (await svc().confirm_file(session["id"], preview["preview_id"], "xlsx-1"))[
        "accepted"
    ] == 1


@pytest.mark.asyncio
async def test_process_distinguishes_instrument_and_product(domain):
    db, svc, ids = domain
    d, s, session, m = await session_setup(domain)
    await svc().ingest(
        MeasurementIngest(session_id=session["id"], source_key="batch-1", measurements=[m])
    )
    record = await svc().get("inspection-sessions", session["id"])
    output = await AGENTS["laboratory_testing"].run(await svc().snapshot(record))
    assert output["early_warning"] and output["status"] == "awaiting_review"
    assert output["probability"] is None and not output["can_make_final_verdict"]
    await svc("platform_operator").device_status(d["id"], {"calibration_status": "expired"})
    output = await AGENTS["laboratory_testing"].run(await svc().snapshot(record))
    assert any(f["type"] == "instrument_suspected" for f in output["findings"])
    assert not output["early_warning"] and output["status"] == "insufficient_evidence"


@pytest.mark.asyncio
async def test_run_persistent_checkpoint_idempotence_and_staleness(domain, monkeypatch):
    db, svc, ids = domain
    d, s, session, m = await session_setup(domain)
    await svc().ingest(
        MeasurementIngest(session_id=session["id"], source_key="batch-1", measurements=[m])
    )
    r = await svc().get("inspection-sessions", session["id"])
    req = RunRequest(version=r.version, request_key="run-1", operation="laboratory")
    run = await svc().queue_run("inspection-sessions", r.id, req)
    assert (await svc().queue_run("inspection-sessions", r.id, req))["id"] == run["id"]
    output = await svc().execute_run(run["id"])
    assert output["status"] == "completed" and output["iteration"] == 1
    assert (await svc().execute_run(run["id"]))["iteration"] == 1
    r = await svc().get("inspection-sessions", r.id)
    run2 = await svc().queue_run(
        "inspection-sessions",
        r.id,
        RunRequest(version=r.version, request_key="run-2", operation="laboratory"),
    )
    await svc().update(
        "inspection-sessions", r.id, RecordUpdate(version=r.version, name="更新后的检测")
    )
    assert (await svc().execute_run(run2["id"]))["status"] == "stale"


@pytest.mark.asyncio
async def test_result_review_then_signoff_and_lock(domain):
    db, svc, ids = domain
    d, s, session, m = await session_setup(domain)
    await svc().ingest(
        MeasurementIngest(session_id=session["id"], source_key="batch-1", measurements=[m])
    )
    record = await svc().get("inspection-sessions", session["id"])
    run = await svc().queue_run(
        "inspection-sessions",
        record.id,
        RunRequest(version=record.version, request_key="laboratory", operation="laboratory"),
    )
    await svc().execute_run(run["id"])
    record = await svc().get("inspection-sessions", record.id)
    with pytest.raises(ValidationError):
        await svc().submit_review(
            "inspection-sessions",
            record.id,
            ReviewRequest(version=record.version, operation="result.signoff"),
        )
    review = await svc().submit_review(
        "inspection-sessions",
        record.id,
        ReviewRequest(version=record.version, operation="result.review"),
    )
    await svc("expert").review_decision(
        review["id"], ReviewDecision(decision="accept", comment="核对完整测量及标准")
    )
    sign = await svc().submit_review(
        "inspection-sessions",
        record.id,
        ReviewRequest(version=record.version, operation="result.signoff"),
    )
    await svc("expert2").review_decision(
        sign["id"], ReviewDecision(decision="accept", comment="签发不合格结果")
    )
    assert record.status == "signed"
    from app.models.result import InspectionResult

    result = await db.scalar(select(InspectionResult))
    assert result.verdict == "fail" and result.reviewed_by == ids["actors"]["expert2"]
    from app.schemas.qdl import parse_qdl

    assert parse_qdl(record.data["qdl_candidate"]).validation_status == "valid"
    with pytest.raises(ConflictError):
        await svc().ingest(
            MeasurementIngest(session_id=record.id, source_key="after-sign", measurements=[m])
        )


@pytest.mark.asyncio
async def test_risk_requires_observed_evidence_and_independent_review():
    snapshot = {
        "standard": {
            "applicability": {"clause_version": "v1", "effective_from": "2026-01-01"},
            "clauses": [{"text": "条款"}],
        },
        "data": {
            "description": "续航下降",
            "enterprise_id": uid(),
            "product_sku_id": uid(),
            "inspection_standard_id": uid(),
            "evidence": [
                {
                    "evidence_id": "real",
                    "source_id": "complaint-1",
                    "source_type": "complaint",
                    "nature": "observed",
                    "text": "续航下降",
                }
            ],
        },
    }
    calls = []

    async def model(name, data):
        calls.append(name)
        return (
            {"summary": "需要检测确认", "risk_level": "medium", "evidence_ids": ["real"]}
            if name == "risk_case.assess"
            else {"supported": True, "issues": []}
        )

    result = await AGENTS["public_opinion_monitoring"].run(snapshot, model)
    assert calls == ["risk_case.assess", "trust.review"] and result["probability"] is None
    assert result["status"] == "awaiting_review"

    async def hallucinate(name, data):
        return {"risk_level": "critical", "evidence_ids": ["invented"]}

    assert (await AGENTS["public_opinion_monitoring"].run(snapshot, hallucinate))[
        "status"
    ] == "manual_review_required"


@pytest.mark.asyncio
async def test_trend_without_exposure_does_not_invent_failure_rate():
    out = await AGENTS["market_monitoring"].run(
        {
            "as_of": "2026-09-15T00:00:00",
            "cases": [
                {"data": {"occurred_at": "2026-09-14T00:00:00", "complaint_region_id": "重庆"}}
            ],
        }
    )
    assert out["current_count"] == 1 and out["probability"] is None


@pytest.mark.asyncio
async def test_sampling_reports_infeasible_coverage():
    out = await AGENTS["supervision_sampling"].run(
        {
            "data": {
                "candidates": [],
                "budget": 10,
                "max_samples": 1,
                "max_staff_hours": 1,
                "hours_per_sample": 1,
                "required_categories": ["电动自行车"],
            },
            "cases": [],
            "devices": [],
        }
    )
    assert out["status"] == "infeasible" and out["uncovered_categories"] == ["电动自行车"]


@pytest.mark.asyncio
async def test_public_api_authorization_and_device_token_rotation(domain):
    from fastapi import FastAPI
    from httpx import AsyncClient, ASGITransport
    from app.api.v1 import supervision as api
    from app.api.v1.deps import get_db, get_current_user

    db, svc, ids = domain
    role = {"value": "admin"}
    app = FastAPI()
    from app.core.error_handlers import register_error_handlers

    register_error_handlers(app)
    app.include_router(api.router, prefix="/api/v1")

    async def session():
        try:
            yield db
            await db.commit()
        except Exception:
            await db.rollback()
            raise

    def current():
        return svc(role["value"]).current

    app.dependency_overrides[get_db] = session
    app.dependency_overrides[get_current_user] = current
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        region = (
            await client.post("/api/v1/regions", json={"code": "重庆", "name": "重庆", "data": {}})
        ).json()["data"]
        assert region["name"] == "重庆"
        role["value"] = "user"
        assert (
            await client.post("/api/v1/regions", json={"code": "越权", "name": "越权", "data": {}})
        ).status_code == 403
        assert (await client.get("/api/v1/regions/not-a-uuid")).status_code == 422
        d, s, session_record, m = await session_setup(domain)
        await db.commit()
        role["value"] = "app_developer"
        token = (await client.post(f"/api/v1/devices/{d['id']}/connection")).json()["data"]["token"]
        req = {"session_id": session_record["id"], "source_key": "http-1", "measurements": [m]}
        result = await client.post(
            "/api/v1/device-ingest/measurements", json=req, headers={"X-Device-Token": token}
        )
        assert result.status_code == 200 and result.json()["data"]["accepted"] == 1
        await client.post(f"/api/v1/devices/{d['id']}/connection")
        assert (
            await client.post(
                "/api/v1/device-ingest/measurements", json=req, headers={"X-Device-Token": token}
            )
        ).status_code == 403
        role["value"] = "algorithm_engineer"
        assert (await client.get("/api/v1/risk-cases")).status_code == 403


@pytest.mark.asyncio
async def test_dispatcher_parallel_overlap_and_state_isolation(monkeypatch):
    from agent.router.manager_dispatcher import ManagerDispatcher
    from agent.router.manager_state import ManagerState
    from agent.router.contracts import AgentPlanStep, AgentRoutePlan
    from agent.contracts.quality_contracts import NormalizedRequest
    from agent.router.executors.base import artifact, observation
    from app.services.task_blackboard_service import TaskBlackboardService
    from app.services.agent_local_memory_service import AgentLocalMemoryService

    blackboard = TaskBlackboardService(redis_url="redis://127.0.0.1:1/0")
    dispatcher = ManagerDispatcher(blackboard=blackboard, local_memory=AgentLocalMemoryService())
    entered = set()
    both = asyncio.Event()
    states = []

    class Executor:
        async def execute(self, step, state, request, db_session=None):
            states.append(state)
            entered.add(step.owner_agent)
            if len(entered) == 2:
                both.set()
            await asyncio.wait_for(both.wait(), timeout=3)
            assert state.current_owner_agent == step.owner_agent
            state.used_llm_calls += 1
            a = artifact(step, "result", content={"owner": step.owner_agent}, summary="完成")
            return observation(
                step, status="success", summary="完成", artifact_ids=[a.artifact_id]
            ), [a]

    dispatcher._executors["vision"] = Executor()
    dispatcher._executors["lab_detection"] = Executor()
    state = ManagerState(
        request_id="parallel", workflow_run_id="parallel", org_id=uid(), original_query="检测"
    )
    request = NormalizedRequest(
        request_id="parallel", org_id=state.org_id, user_id=uid(), query="检测"
    )
    plan = AgentRoutePlan(
        plan_id="P",
        surface="chat",
        goal="并行检查",
        steps=[
            AgentPlanStep(
                step_id="s1",
                owner_agent="vision",
                capability="vision.inspect",
                parallel_group="professional",
            ),
            AgentPlanStep(
                step_id="s2",
                owner_agent="lab_detection",
                capability="lab.early_risk.assess",
                parallel_group="professional",
            ),
        ],
    )
    observations, artifacts = await dispatcher.dispatch(plan, state, request)
    assert len(entered) == 2 and len(artifacts) == 2 and state.used_llm_calls == 2
    assert states[0] is not states[1] and states[0] is not state
    assert len(state.artifacts) == 2


@pytest.mark.asyncio
async def test_image_draft_does_not_create_formal_result(domain):
    from agent.contracts.quality_contracts import AgentOutput
    from app.services.inspection_pipeline_service import _save_supervision_image_draft
    from app.models.result import InspectionResult

    db, svc, ids = domain
    task = await db.get(InspectionTask, ids["task"])
    out = AgentOutput(
        answer="可观察部位存在破损",
        summary="图片草稿",
        result_card={"verdict": "fail"},
        citations=[{"source_id": "条款"}],
    )
    draft = await _save_supervision_image_draft(db, task, out)
    assert draft.status == "awaiting_review"
    assert await db.scalar(select(InspectionResult)) is None


@pytest.mark.asyncio
async def test_invalid_typed_measurement_keeps_pending_raw_data(domain):
    db, svc, ids = domain
    d, s, session, m = await session_setup(domain)
    req = MeasurementIngest(
        session_id=session["id"], source_key="pending-1", measurements=[{**m, "unit": "unknown"}]
    )
    result = await svc().ingest_or_pending(req)
    assert result["pending"] == 1 and result["accepted"] == 0
    assert (await svc().ingest_or_pending(req))["pending"] == 1
    assert not list(await db.scalars(select(MeasurementRecord)))


@pytest.mark.asyncio
async def test_region_dictionary_import_resolves_parent_codes_and_rejects_cycle(domain):
    from app.services.supervision_imports import preview, confirm

    db, svc, ids = domain
    raw = "code,name,parent_code,dictionary_version\nCQ-YB,渝北,CQ,2026\nCQ,重庆,,2026\n".encode()
    result = await preview(svc("admin"), "regions", "regions.csv", raw, {})
    assert result["can_confirm"]
    imported = await confirm(svc("admin"), "regions", result["preview_id"])
    assert imported["accepted"] == 2
    rows = (await svc("admin").list("regions"))["items"]
    parent = next(r for r in rows if r["code"] == "CQ")
    child = next(r for r in rows if r["code"] == "CQ-YB")
    assert child["data"]["parent_id"] == parent["id"]
    assert (await confirm(svc("admin"), "regions", result["preview_id"]))["accepted"] == 2
    bad = await preview(
        svc("admin"), "regions", "cycle.csv", "code,name,parent_code\nA,A,B\nB,B,A\n".encode(), {}
    )
    assert not bad["can_confirm"] and bad["errors"]


@pytest.mark.asyncio
async def test_signed_result_cannot_be_changed_through_legacy_review_or_task_delete(domain):
    from app.models.result import InspectionResult
    from app.services.result_service import ResultService
    from app.services.task_service import TaskService

    db, svc, ids = domain
    task = await db.get(InspectionTask, ids["task"])
    task.meta_data = {"supervision": True, "input_mode": "measurement"}
    task.status = "done"
    result = InspectionResult(
        id=uid(),
        org_id=ids["org"],
        task_id=task.id,
        verdict="fail",
        overall_score=0,
        defects=[],
        reasoning_chain={"supervision_signed": True},
        llm_model="rules",
        prompt_version="v1",
        created_at=datetime.now(),
        updated_at=datetime.now(),
    )
    db.add(result)
    await db.flush()
    with pytest.raises(ConflictError):
        await ResultService(db, ids["org"]).review(
            result.id, ids["actors"]["expert"], "expert", {"verdict": "pass"}
        )
    with pytest.raises(ConflictError):
        await TaskService(
            db, ids["org"], actor_user_id=ids["actors"]["user"], actor_role="user"
        ).delete_task(task.id)
    assert await db.get(InspectionResult, result.id) is not None


@pytest.mark.asyncio
async def test_algorithm_only_receives_explicitly_enrolled_structured_result(domain):
    from app.models.dataset import Dataset
    from app.models.result import InspectionResult
    from app.services.task_result_ingest_service import TaskResultIngestService
    from app.schemas.task import TaskResultIngestRequest

    db, svc, ids = domain
    task = await db.get(InspectionTask, ids["task"])
    task.meta_data = {"input_mode": "measurement", "supervision": True}
    task.status = "done"
    now = datetime.now()
    dataset_id = uid()
    db.add(
        Dataset(
            id=dataset_id,
            org_id=ids["org"],
            created_by=ids["actors"]["algorithm_engineer"],
            name="已授权测量集",
            status="active",
            modality="text",
            created_at=now,
            updated_at=now,
        )
    )
    db.add(
        InspectionResult(
            id=uid(),
            org_id=ids["org"],
            task_id=task.id,
            verdict="fail",
            overall_score=0,
            defects=[],
            reasoning_chain={
                "supervision_signed": True,
                "findings": [{"item": "温升", "value": 55}],
            },
            llm_model="rules",
            prompt_version="v1",
            created_at=now,
            updated_at=now,
        )
    )
    session = SupervisionRecord(
        org_id=ids["org"],
        kind="inspection-sessions",
        code="signed-text",
        name="检测结果",
        created_by=ids["actors"]["user"],
        status="signed",
        data={"task_id": task.id, "knowledge_status": "current"},
    )
    db.add(session)
    await db.flush()
    ingest = TaskResultIngestService(
        db,
        ids["org"],
        actor_user_id=ids["actors"]["algorithm_engineer"],
        actor_role="algorithm_engineer",
    )
    payload = TaskResultIngestRequest(target="dataset", dataset_id=dataset_id, mode="candidate")
    with pytest.raises(ForbiddenError):
        await ingest.ingest_task_result(task_id=task.id, payload=payload)
    await svc().enroll_dataset(session.id, dataset_id)
    result = await ingest.ingest_task_result(task_id=task.id, payload=payload)
    assert result.created_sample_count == 1
    assert (await ingest.ingest_task_result(task_id=task.id, payload=payload)).skipped_count == 1


def test_calibration_and_response_time_metrics():
    from app.services.supervision_evaluation import calibration_report, response_time_report

    report = calibration_report([0, 1], [0, 1])
    assert report["brier"] == 0 and report["ece"] == 0
    assert calibration_report([0.5, 0.5], [0, 1])["brier"] == 0.25
    assert response_time_report([100, 100], [40, 40])["median_reduction"] == 0.6


def test_supervision_migration_upgrades_and_preserves_legacy_tables():
    import importlib.util
    from pathlib import Path
    from sqlalchemy import create_engine, inspect, text
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    path = Path(__file__).resolve().parents[1] / "migrations/versions/0100_quality_supervision.py"
    spec = importlib.util.spec_from_file_location("supervision_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        for table in ("product_skus", "product_batches", "inspection_standard_libraries"):
            connection.execute(text(f"CREATE TABLE {table} (id INTEGER PRIMARY KEY, name TEXT)"))
            connection.execute(text(f"INSERT INTO {table} (id,name) VALUES (1,'legacy')"))
        with Operations.context(MigrationContext.configure(connection)):
            module.upgrade()
            assert set(module.TABLES).issubset(inspect(connection).get_table_names())
            for table in module.TABLES:
                assert {c["name"] for c in inspect(connection).get_columns(table)}.issubset(
                    {c.name for c in Base.metadata.tables[table].c}
                )
            module.downgrade()
        assert set(inspect(connection).get_table_names()) == {
            "product_skus",
            "product_batches",
            "inspection_standard_libraries",
        }
        assert connection.scalar(text("SELECT name FROM product_skus WHERE id=1")) == "legacy"
    engine.dispose()


def test_supervision_exposes_exactly_the_four_declared_business_agents():
    assert set(AGENTS) == {
        "market_monitoring",
        "public_opinion_monitoring",
        "supervision_sampling",
        "laboratory_testing",
    }


def test_risk_monitoring_and_situation_are_separate_manager_capabilities():
    from app.services.supervision_service import AGENT_CAPABILITIES, AGENT_OPERATIONS

    assert AGENT_OPERATIONS["risk_monitoring"] == (
        "risk-cases",
        "public_opinion_monitoring",
    )
    assert "public_opinion" not in AGENT_OPERATIONS
    assert "market_monitor" not in AGENT_OPERATIONS
    assert AGENT_CAPABILITIES["public_opinion_monitoring"] == (
        "risk_case.assess",
        "risk_case_assessment",
    )
    assert AGENT_CAPABILITIES["market_monitoring"] == (
        "risk_situation.analyze",
        "risk_situation_report",
    )


def test_adaptive_plan_items_preserve_legacy_test_items_and_reject_overlap():
    from pydantic import ValidationError as SchemaValidationError

    from app.schemas.supervision import Candidate, SessionData

    item = {"item": "温升", "unit": "C", "method": "GB/T-test"}
    legacy = Candidate(case_id=uid(), test_items=[item])
    assert len(legacy.required_items) == 1
    assert legacy.candidate_items == []

    adaptive = Candidate(case_id=uid(), required_items=[item], candidate_items=[])
    assert adaptive.test_items == adaptive.required_items

    session = SessionData(
        task_id=uid(),
        required_items=[item],
        candidate_items=[{"item": "续航", "unit": "km", "method": "road-test"}],
        adaptive_enabled=True,
    )
    assert [entry.item for entry in session.test_items] == ["温升", "续航"]

    with pytest.raises(SchemaValidationError):
        Candidate(case_id=uid(), required_items=[item], candidate_items=[item])


def test_inspection_goal_contract_requires_items_and_strong_stop_policy():
    from pydantic import ValidationError as SchemaValidationError

    from app.schemas.supervision import InspectionGoalCreate

    payload = {
        "session_version": 3,
        "request_key": "goal:session:3",
        "risk_hypotheses": ["温升异常"],
        "required_items": [{"item": "温升", "unit": "C", "method": "GB/T-test"}],
        "candidate_items": [{"item": "续航", "unit": "km", "method": "road-test"}],
        "success_criteria": {"resolve_all_hypotheses": True},
        "stop_policy": {
            "rule_version": "adaptive-stop-v1",
            "evidence_sufficiency_threshold": 0.85,
            "require_expert_approval": True,
        },
    }
    goal = InspectionGoalCreate.model_validate(payload)
    assert goal.stop_policy.require_expert_approval is True
    assert len(goal.required_items) == 1 and len(goal.candidate_items) == 1

    with pytest.raises(SchemaValidationError):
        InspectionGoalCreate.model_validate(
            {**payload, "required_items": [], "candidate_items": []}
        )

    with pytest.raises(SchemaValidationError):
        InspectionGoalCreate.model_validate(
            {**payload, "candidate_items": list(payload["required_items"])}
        )

    with pytest.raises(SchemaValidationError):
        InspectionGoalCreate.model_validate(
            {
                **payload,
                "stop_policy": {
                    "rule_version": "adaptive-stop-v1",
                    "evidence_sufficiency_threshold": 1.1,
                },
            }
        )


@pytest.mark.asyncio
async def test_adaptive_inspection_goal_next_test_and_strong_stop_gate(domain):
    from app.schemas.supervision import (
        InspectionGoalCreate,
        NextTestDecisionCreate,
        StopDecisionCreate,
    )
    from app.services.adaptive_inspection_service import AdaptiveInspectionService

    db, svc, _ids = domain
    _device, _sample, session, measurement = await session_setup(domain)
    adaptive = AdaptiveInspectionService(db, svc().current)
    goal_payload = InspectionGoalCreate(
        session_version=session["version"],
        request_key=f"goal:{session['id']}:{session['version']}",
        risk_hypotheses=["温升异常"],
        required_items=[
            {
                "item": "温升",
                "unit": "C",
                "method": "试点方法",
                "upper_limit": 45,
                "standard_ref": "试点条款1",
            }
        ],
        candidate_items=[],
        success_criteria={"resolve_all_hypotheses": True},
        stop_policy={
            "rule_version": "adaptive-stop-v1",
            "evidence_sufficiency_threshold": 0.8,
            "require_expert_approval": True,
        },
    )
    goal = await adaptive.create_goal(session["id"], goal_payload)
    assert goal["status"] == "active" and len(goal["required_items"]) == 1
    assert (await adaptive.create_goal(session["id"], goal_payload))["id"] == goal["id"]

    evidence = await adaptive.get_latest_evidence_state(session["id"])
    next_test = await adaptive.propose_next_test(
        session["id"],
        NextTestDecisionCreate(
            evidence_state_id=evidence["id"],
            request_key="next-test:round-0",
            rule_version="next-test-v1",
        ),
    )
    assert next_test["kind"] == "next_test"
    assert next_test["payload"]["selected"]["item_class"] == "required"

    await svc().ingest(
        MeasurementIngest(
            session_id=session["id"],
            measurements=[measurement],
            source_key="adaptive-1",
        )
    )
    evidence = await adaptive.get_latest_evidence_state(session["id"])
    assert evidence["round"] == 1
    assert evidence["device_trust_state"]["status"] == "trusted"
    assert evidence["product_quality_state"]["status"] == "abnormal"
    assert evidence["evidence_sufficiency"] == 1.0

    stop = await adaptive.propose_stop(
        session["id"],
        StopDecisionCreate(
            evidence_state_id=evidence["id"],
            request_key="stop:round-0",
            rule_version="adaptive-stop-v1",
        ),
    )
    assert stop["kind"] == "goal_reached"
    assert stop["status"] == "awaiting_approval"
    with pytest.raises(ForbiddenError):
        await adaptive.review_stop_decision(
            session["id"],
            stop["id"],
            approve=True,
            comment="本人不能审批",
        )
    reviewed = await AdaptiveInspectionService(
        db, svc("expert").current
    ).review_stop_decision(
        session["id"],
        stop["id"],
        approve=True,
        comment="必检项完成且证据充分，同意结束本轮采集",
    )
    assert reviewed["status"] == "approved"
    assert (await svc().get("inspection-sessions", session["id"])).status == "goal_reached"
    events = list(
        await db.scalars(
            select(SupervisionEvent).where(
                SupervisionEvent.org_id == svc().org_id,
                SupervisionEvent.aggregate_id == session["id"],
            )
        )
    )
    assert {event.event_type for event in events} >= {
        "inspection.goal.created",
        "inspection.next_test.proposed",
        "inspection.stop.proposed",
        "inspection.stop.approved",
    }


def test_adaptive_supervision_migration_adds_and_removes_schema():
    import importlib.util
    from pathlib import Path

    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import create_engine, inspect, text

    path = (
        Path(__file__).resolve().parents[1]
        / "migrations/versions/0102_adaptive_quality_supervision.py"
    )
    spec = importlib.util.spec_from_file_location("adaptive_supervision_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    engine = create_engine("sqlite:///:memory:")
    added_tables = {
        "inspection_goals",
        "inspection_evidence_states",
        "inspection_decisions",
        "device_commands",
        "supervision_events",
    }
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE measurement_records (id BLOB PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE inspection_results (id BLOB PRIMARY KEY)"))
        with Operations.context(MigrationContext.configure(connection)):
            module.upgrade()
            inspector = inspect(connection)
            assert added_tables.issubset(inspector.get_table_names())
            assert "request_key" in {
                column["name"] for column in inspector.get_columns("inspection_goals")
            }
            assert {
                "command_id",
                "sequence_no",
                "validation_status",
                "calibration_version",
            }.issubset({column["name"] for column in inspector.get_columns("measurement_records")})
            assert {
                "inspection_session_id",
                "stop_decision_id",
                "result_status",
                "signed_by",
                "signed_at",
            }.issubset({column["name"] for column in inspector.get_columns("inspection_results")})
            module.downgrade()
        assert set(inspect(connection).get_table_names()) == {
            "measurement_records",
            "inspection_results",
        }
    engine.dispose()
