"""Raw intake: optional origins, consistent file imports, and object-scoped access."""
import csv
import hashlib
import io
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import String, exists, func, or_, select

from app.core.exceptions import ConflictError, NotFoundError, ValidationError
from app.core.permissions import has_permission, require_role
from app.models.quality_risk import (
    QualityAttachment,
    QualityDataSource,
    QualityIngestionJob,
    QualityRiskCase,
    QualitySourceRecord,
)
from app.models.supervision import SupervisionRecord
from app.schemas.quality_risk import RECORD_TYPES, QualityEventCreate
from app.services.business_number_service import BusinessNumberService


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str, allow_nan=False).encode()).hexdigest()


def cleaned(value):
    return "" if value is None else str(value).strip()


def json_cell(value, default):
    if value is None or value == "":
        return default
    return json.loads(value) if isinstance(value, str) else value


class QualityIntakeMixin:
    def require_intake(self, write=False):
        if write:
            require_role("quality_material_submit", self.current.role)
        elif not has_permission("quality_material_read_business", self.current.role):
            require_role("quality_material_read_own", self.current.role)

    def readable_records(self):
        stmt = select(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.deleted_at.is_(None))
        if not has_permission("quality_material_read_business", self.current.role):
            stmt = stmt.where(QualitySourceRecord.created_by == self.actor)
        return stmt

    async def scoped_records(self):
        if has_permission("quality_material_read_business", self.current.role):
            return self.readable_records()
        cases = await self.db.scalars(select(QualityRiskCase.source_record_ids).where(
            QualityRiskCase.org_id == self.org_id, QualityRiskCase.assigned_to == self.actor,
            QualityRiskCase.deleted_at.is_(None),
        ))
        assigned_ids = {str(record_id) for ids in cases for record_id in (ids or [])}
        return select(QualitySourceRecord).where(
            QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.deleted_at.is_(None),
            or_(QualitySourceRecord.created_by == self.actor, QualitySourceRecord.id.in_(assigned_ids)),
        )

    async def intake_source(self, source_id, method="manual"):
        if source_id:
            row = await self.db.scalar(select(QualityDataSource).where(
                QualityDataSource.id == source_id, QualityDataSource.org_id == self.org_id,
                QualityDataSource.deleted_at.is_(None),
            ).with_for_update())
            if not row:
                raise ValidationError("数据来源不可用")
        else:
            code = f"BUILTIN-{method.upper()}"
            values = {"org_id": self.org_id, "code": code, "name": "文件资料接入" if method == "file" else "人工资料接入", "source_type": "mixed", "connector_type": method, "config": {"builtin": True}, "status": "active", "created_by": self.actor}
            from app.core.ids import uuid7
            values["id"] = str(uuid7())
            if self.db.get_bind().dialect.name == "mysql":
                from sqlalchemy.dialects.mysql import insert
                statement = insert(QualityDataSource).values(**values)
                statement = statement.on_duplicate_key_update(code=statement.inserted.code)
            else:
                from sqlalchemy.dialects.sqlite import insert
                statement = insert(QualityDataSource).values(**values).on_conflict_do_nothing(index_elements=["org_id", "code"])
            await self.db.execute(statement)
            row = await self.db.scalar(select(QualityDataSource).where(
                QualityDataSource.org_id == self.org_id, QualityDataSource.code == code,
            ).with_for_update())
            await BusinessNumberService(self.db).assign(row)
        if row.status != "active" or row.deleted_at is not None:
            raise ValidationError("数据来源未启用")
        return row

    async def validate_intake(self, payload):
        if payload.record_type == "standard_document":
            raise ValidationError("标准文件请通过标准库上传，业务资料接入不重复登记标准")
        try:
            digest(payload.content)
        except ValueError as exc:
            raise ValidationError("资料包含非法数值") from exc
        if not payload.attachment_ids and not any(v not in (None, "", [], {}) for v in payload.content.values()):
            raise ValidationError("请提供原始内容或附件")
        if payload.product_ref.category_id:
            await self.category(payload.product_ref.category_id)
        if payload.product_ref.product_id:
            product = await self.product(payload.product_ref.product_id)
            if payload.product_ref.category_id and product.category_id != payload.product_ref.category_id:
                raise ValidationError("产品不属于所选产品类别")
        if payload.product_ref.production_batch_id:
            if not payload.product_ref.product_id:
                raise ValidationError("关联批次前请先指定所属产品")
            batch = await self.production_batch(payload.product_ref.product_id, payload.product_ref.production_batch_id)
            if batch.original_batch_no and payload.product_ref.production_batch_ref and batch.original_batch_no != payload.product_ref.production_batch_ref:
                raise ValidationError("原始批次号与所选批次不一致")
        if payload.enterprise_ref.enterprise_id:
            enterprise = await self.db.scalar(select(SupervisionRecord.id).where(
                SupervisionRecord.id == payload.enterprise_ref.enterprise_id, SupervisionRecord.org_id == self.org_id,
                SupervisionRecord.kind == "enterprises", SupervisionRecord.status != "archived",
            ).with_for_update(read=True))
            if not enterprise:
                raise ValidationError("企业不属于当前组织或不可用")
        await self.validate_attachments(payload.attachment_ids)

    async def create_event(self, payload: QualityEventCreate, *, import_job_id=None):
        self.require_intake(write=True)
        if payload.ingestion_job_id and payload.ingestion_job_id != import_job_id:
            raise ValidationError("导入任务关联由系统生成，请通过导入确认提交")
        source = await self.intake_source(payload.source_id, "file" if payload.ingestion_job_id else "manual")
        await self.validate_intake(payload)
        data = payload.model_dump(mode="json")
        # Preserve the fingerprint of older submissions with no batch registry reference.
        if not payload.product_ref.production_batch_id:
            data["product_ref"].pop("production_batch_id", None)
        provenance = dict(payload.provenance)
        if payload.origin_label:
            provenance["origin_label"] = payload.origin_label
        origin = cleaned(provenance.get("origin_label"))
        identity = {k: data[k] for k in ("external_record_id", "record_type", "occurred_at", "content", "enterprise_ref", "product_ref", "location", "attachment_ids", "data_nature")}
        identity["provenance"] = {k: v for k, v in provenance.items() if k not in {"imported", "import_batch_code", "import_row", "original_row", "source_timezone", "channel", "request_key"}}
        # Different private submitters retain independent raw statements.
        if not has_permission("quality_material_read_business", self.current.role):
            identity["private_submitter"] = self.actor
        try:
            content_hash = digest(identity)
        except ValueError as exc:
            raise ValidationError("资料包含非法数值") from exc
        candidates = select(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.source_id == source.id).with_for_update()
        if payload.external_record_id:
            previous = list(await self.db.scalars(candidates.where(QualitySourceRecord.external_record_id == payload.external_record_id)))
            previous = [x for x in previous if cleaned(x.provenance.get("origin_label")) == origin and
                        (has_permission("quality_material_read_business", self.current.role) or x.created_by == self.actor)]
            if previous and all(x.content_hash != content_hash for x in previous):
                raise ConflictError("同一出处的外部编号已存在不同内容，请明确修订资料")
        existing = await self.db.scalar(candidates.where(QualitySourceRecord.content_hash == content_hash))
        if existing:
            if existing.created_by != self.actor and not has_permission("quality_material_read_business", self.current.role):
                await self.record(existing.id)
            await BusinessNumberService(self.db).attach([existing])
            return {**self.record_response(existing), "duplicate": True}
        row = QualitySourceRecord(org_id=self.org_id, source_id=source.id, ingestion_job_id=payload.ingestion_job_id,
                                  external_record_id=payload.external_record_id, record_type=payload.record_type,
                                  occurred_at=payload.occurred_at, content=payload.content, content_hash=content_hash,
                                  enterprise_ref=payload.enterprise_ref.model_dump(mode="json", exclude_none=True),
                                  product_ref=payload.product_ref.model_dump(mode="json", exclude_none=True),
                                  location=payload.location.model_dump(mode="json", exclude_none=True),
                                  attachment_ids=payload.attachment_ids, provenance=provenance,
                                  authorization_scope="private" if not has_permission("quality_material_read_business", self.current.role) else payload.authorization_scope,
                                  data_nature=payload.data_nature, normalization_status="stored", normalization_errors=[], created_by=self.actor)
        self.db.add(row)
        await self.flush_and_refresh(row)
        await self.materialize_evidence(row)
        return {**self.record_response(row), "duplicate": False}

    async def list_records(self, *, page, size, record_type=None, record_types=None, keyword=None, product_id=None,
                           production_batch_ref=None, source_id=None, ingestion_job_id=None):
        self.require_intake()
        stmt = await self.scoped_records()
        if record_type:
            stmt = stmt.where(QualitySourceRecord.record_type == record_type)
        if record_types:
            if set(record_types) - RECORD_TYPES:
                raise ValidationError("包含未知资料类型")
            stmt = stmt.where(QualitySourceRecord.record_type.in_(record_types))
        if keyword:
            stmt = stmt.where(or_(QualitySourceRecord.content.cast(String).contains(keyword, autoescape=True),
                                  QualitySourceRecord.external_record_id.contains(keyword, autoescape=True)))
        if product_id:
            stmt = stmt.where(QualitySourceRecord.product_ref["product_id"].as_string() == product_id)
        if production_batch_ref:
            stmt = stmt.where(QualitySourceRecord.product_ref["production_batch_ref"].as_string() == production_batch_ref)
        if source_id:
            stmt = stmt.where(QualitySourceRecord.source_id == source_id)
        if ingestion_job_id:
            stmt = stmt.where(QualitySourceRecord.ingestion_job_id == ingestion_job_id)
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery()))
        rows = list(await self.db.scalars(stmt.order_by(QualitySourceRecord.received_at.desc(), QualitySourceRecord.id.desc()).offset((page-1)*size).limit(size)))
        await BusinessNumberService(self.db).attach(rows)
        return {"items": [self.record_response(x) for x in rows], "total": int(total or 0), "page": page, "size": size}

    async def get_record(self, record_id):
        self.require_intake()
        return self.record_response(await self.record(record_id))

    async def record(self, record_id):
        row = await self.db.scalar((await self.scoped_records()).where(QualitySourceRecord.id == record_id))
        if not row:
            raise NotFoundError("来源记录不存在或无权访问")
        await BusinessNumberService(self.db).attach([row])
        return row

    async def accessible_attachment(self, attachment_id):
        row = await self.db.scalar(select(QualityAttachment).where(QualityAttachment.id == attachment_id,
                                 QualityAttachment.org_id == self.org_id, QualityAttachment.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("附件不存在或无权访问")
        if row.uploaded_by != self.actor and not has_permission("quality_material_read_business", self.current.role):
            related = (await self.scoped_records()).where(QualitySourceRecord.attachment_ids.cast(String).contains(f'"{attachment_id}"', autoescape=True))
            if not await self.db.scalar(select(exists(related))):
                raise NotFoundError("附件不存在或无权访问")
        return row

    async def validate_attachments(self, ids):
        for attachment_id in set(ids):
            await self.accessible_attachment(attachment_id)

    async def upload_attachment(self, file):
        self.require_intake(write=True)
        data = await file.read(50*1024*1024+1)
        if not data or len(data) > 50*1024*1024:
            raise ValidationError("附件不能为空且不能超过50MB")
        name = Path(file.filename or "").name
        allowed = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".txt", ".csv", ".xlsx", ".json"}
        if Path(name).suffix.lower() not in allowed:
            raise ValidationError("首版附件支持PDF、图片、文本、CSV、XLSX和JSON")
        checksum = hashlib.sha256(data).hexdigest()
        # Include the uploader so two private uploads never share a public object key.
        key = f"quality-data/{self.org_id}/{self.actor}/{checksum}-{name}"
        stored = self.storage.put_bytes(bucket="quality-data", object_key=key, data=data, content_type=file.content_type)
        row = QualityAttachment(org_id=self.org_id, file_name=name, mime_type=file.content_type or "application/octet-stream",
                                size_bytes=len(data), sha256=checksum, bucket=str(stored.get("bucket") or "quality-data"),
                                object_key=str(stored.get("object_key") or key), uploaded_by=self.actor)
        self.db.add(row)
        await self.flush_and_refresh(row)
        return self.attachment_response(row)

    async def preview_import(self, *, source_id=None, request_key, file, mapping=None, origin_label=None, record_type=None):
        self.require_intake(write=True)
        require_role("quality_ingestion_manage", self.current.role)
        mapping = mapping or {}
        content = await file.read(20*1024*1024+1)
        if not content or len(content) > 20*1024*1024:
            raise ValidationError("导入文件不能为空且不能超过20MB")
        rows = self.parse_rows(file.filename or "", content)
        columns = list(dict.fromkeys(k for row in rows for k in row if k != "__intake_row__"))
        aliases = {"text": ["原始内容", "投诉内容", "内容", "描述"], "record_type": ["资料类型", "数据类型"], "occurred_at": ["发生时间", "事件时间"], "origin_label": ["原始出处", "来源"], "external_record_id": ["外部编号", "报告编号", "投诉编号"], "enterprise_name": ["企业名称"], "enterprise_id": ["企业ID"], "enterprise_role": ["企业角色"], "product_id": ["产品ID"], "category_id": ["类别ID"], "production_batch_ref": ["生产批次", "批次号"], "unit_serial_ref": ["单件序列号"], "formatted_address": ["地点", "地址"]}
        suggested = {target: next((column for column in options if column in columns), "") for target, options in aliases.items() if target not in columns}
        mapping = {**{k: v for k, v in suggested.items() if v}, **{k: v for k, v in mapping.items() if v}}
        fingerprint = digest({"file_hash": hashlib.sha256(content).hexdigest(), "file_name": file.filename,
                              "source_id": source_id, "mapping": mapping, "origin_label": origin_label, "record_type": record_type})
        existing = await self.db.scalar(select(QualityIngestionJob).where(QualityIngestionJob.org_id == self.org_id, QualityIngestionJob.request_key == request_key))
        if existing:
            if existing.created_by != self.actor and not has_permission("quality_material_read_business", self.current.role):
                raise NotFoundError("导入任务不存在或无权访问")
            await BusinessNumberService(self.db).attach([existing])
            if existing.raw_payload.get("fingerprint") != fingerprint:
                raise ConflictError("同一请求键对应的文件、来源或映射已变化")
            return self.ingestion_response(existing)
        if not rows or len(rows) > 5000:
            raise ValidationError("导入需包含1至5000条记录，超限请拆分文件")
        source = await self.intake_source(source_id, "file")
        existing = await self.db.scalar(select(QualityIngestionJob).where(QualityIngestionJob.org_id == self.org_id, QualityIngestionJob.request_key == request_key).with_for_update())
        if existing:
            if existing.created_by != self.actor and not has_permission("quality_material_read_business", self.current.role):
                raise NotFoundError("导入任务不存在或无权访问")
            await BusinessNumberService(self.db).attach([existing])
            if existing.raw_payload.get("fingerprint") != fingerprint:
                raise ConflictError("同一请求键对应的文件、来源或映射已变化")
            return self.ingestion_response(existing)
        normalized, errors = [], []
        for index, row in enumerate(rows, 1):
            row = dict(row)
            location = row.pop("__intake_row__", index)
            try:
                mapped = dict(row)
                mapped.update({target: row.get(column) for target, column in mapping.items()})
                if not cleaned(mapped.get("record_type")) and record_type:
                    mapped["record_type"] = record_type
                raw = self.row_to_event(mapped, source.id)
                raw["origin_label"] = cleaned(mapped.get("origin_label")) or origin_label
                raw["provenance"] = {**raw["provenance"], "import_row": location, "original_row": row}
                payload = QualityEventCreate.model_validate(raw)
                await self.validate_intake(payload)
                normalized.append(payload.model_dump(mode="json"))
            except (TypeError, ValueError, ValidationError, NotFoundError) as exc:
                errors.append({"row": location, "message": str(exc)})
        number = await BusinessNumberService(self.db).allocate(self.org_id, "DR")
        job = QualityIngestionJob(org_id=self.org_id, source_id=source.id, request_key=request_key, import_batch_code=number,
                                  status="preview", file_name=file.filename, mapping=mapping,
                                  raw_payload={"rows": normalized, "fingerprint": fingerprint, "source_columns": columns}, total_count=len(rows),
                                  valid_count=len(normalized), error_count=len(errors), errors=errors, created_by=self.actor)
        self.db.add(job)
        await self.db.flush()
        await BusinessNumberService(self.db).assign(job, reserved=number)
        await self.db.refresh(job)
        return self.ingestion_response(job)

    async def confirm_import(self, job_id):
        self.require_intake(write=True)
        require_role("quality_ingestion_manage", self.current.role)
        job = await self.ingestion_job(job_id, lock=True)
        if job.status == "confirmed":
            return self.ingestion_response(job)
        if job.status != "preview" or job.error_count:
            raise ConflictError("请修复预览错误后确认；已取消的导入不能确认")
        outcome = {"created_count": 0, "duplicate_count": 0, "failed_count": 0, "records": []}
        for raw in job.raw_payload.get("rows", []):
            payload = QualityEventCreate.model_validate({**raw, "ingestion_job_id": job.id})
            result = await self.create_event(payload, import_job_id=job.id)
            outcome["duplicate_count" if result["duplicate"] else "created_count"] += 1
            outcome["records"].append({"row": payload.provenance.get("import_row"), "record_id": result["id"], "duplicate": result["duplicate"]})
        job.raw_payload = {**job.raw_payload, "outcome": outcome}
        job.status = "confirmed"
        await self.flush_and_refresh(job)
        return self.ingestion_response(job)

    async def cancel_import(self, job_id):
        self.require_intake(write=True)
        require_role("quality_ingestion_manage", self.current.role)
        job = await self.ingestion_job(job_id, lock=True)
        if job.status == "confirmed":
            raise ConflictError("已确认导入不能取消")
        job.status = "cancelled"
        await self.flush_and_refresh(job)
        return self.ingestion_response(job)

    async def ingestion_job(self, job_id, *, lock=False):
        self.require_intake()
        stmt = select(QualityIngestionJob).where(QualityIngestionJob.id == job_id, QualityIngestionJob.org_id == self.org_id,
                                                QualityIngestionJob.deleted_at.is_(None))
        if not has_permission("quality_material_read_business", self.current.role):
            stmt = stmt.where(QualityIngestionJob.created_by == self.actor)
        row = await self.db.scalar(stmt.with_for_update() if lock else stmt)
        if not row:
            raise NotFoundError("导入任务不存在或无权访问")
        await BusinessNumberService(self.db).attach([row])
        return row

    def parse_rows(self, name, content):
        suffix = Path(name).suffix.lower()
        try:
            if suffix == ".csv":
                reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
                headers = [cleaned(h) for h in (reader.fieldnames or [])]
                if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
                    raise ValidationError("CSV表头不能为空或重复")
                reader.fieldnames = headers
                rows = []
                for row in reader:
                    if None in row:
                        raise ValidationError(f"CSV第{reader.line_num}行列数超过表头")
                    if any(v for v in row.values()):
                        rows.append({**row, "__intake_row__": reader.line_num})
                    if len(rows) > 5000:
                        break
                return rows
            if suffix == ".xlsx":
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                try:
                    values = workbook.active.iter_rows(values_only=True)
                    header = next(values, None)
                    if header is None:
                        return []
                    headers = [cleaned(x) for x in header]
                    while headers and not headers[-1]:
                        headers.pop()
                    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
                        raise ValidationError("XLSX表头不能为空或重复")
                    rows = []
                    for row_index, row in enumerate(values, 2):
                        if any(x is not None for x in row):
                            rows.append({**{headers[i]: row[i] for i in range(min(len(headers), len(row)))}, "__intake_row__": row_index})
                        if len(rows) > 5000:
                            break
                    return rows
                finally:
                    workbook.close()
            if suffix == ".json":
                value = json.loads(content.decode("utf-8-sig"))
                rows = value if isinstance(value, list) else [value]
                if not all(isinstance(r, dict) for r in rows):
                    raise ValueError("JSON每条记录必须是对象")
                return rows
        except (ValueError, UnicodeError, OSError) as exc:
            raise ValidationError(f"无法读取导入文件：{exc}") from exc
        raise ValidationError("结构化导入仅支持CSV、XLSX和JSON；PDF和图片请作为资料附件上传")

    def row_to_event(self, row, source_id):
        labels = {"消费投诉":"consumer_complaint", "监督抽查":"supervision_inspection", "执法案例":"enforcement_case", "企业信息":"enterprise_information", "舆情信息":"public_opinion", "政策文件":"policy_document", "标准文件":"standard_document", "检测报告":"inspection_report", "生产数据":"production_data", "图像":"image", "视频":"video", "设备观测":"device_observation", "专家意见":"expert_opinion"}
        record_type = cleaned(row.get("record_type"))
        occurred_at = row.get("occurred_at") or None
        source_timezone = None
        if occurred_at:
            parsed = datetime.fromisoformat(occurred_at.replace("Z", "+00:00")) if isinstance(occurred_at, str) else occurred_at
            if isinstance(parsed, datetime) and parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=ZoneInfo("Asia/Shanghai"))
                source_timezone = "Asia/Shanghai"
            occurred_at = parsed
        content = row.get("content")
        if isinstance(content, str):
            try:
                content = json.loads(content)
            except json.JSONDecodeError:
                content = {"text": content}
        product = dict(json_cell(row.get("product_ref"), {}))
        enterprise = dict(json_cell(row.get("enterprise_ref"), {}))
        location = dict(json_cell(row.get("location"), {}))
        for key in ("category_id", "product_id", "production_batch_ref", "unit_serial_ref", "source_product_id"):
            if cleaned(row.get(key)):
                product[key] = cleaned(row[key])
        for key, column in (("enterprise_id", "enterprise_id"), ("name", "enterprise_name"), ("credit_code", "credit_code"), ("role", "enterprise_role")):
            if cleaned(row.get(column)):
                enterprise[key] = cleaned(row[column])
        if cleaned(row.get("formatted_address")):
            location.update(formatted_address=cleaned(row["formatted_address"]), location_method="imported")
        return {"source_id": source_id, "external_record_id": cleaned(row.get("external_record_id")) or None,
                "record_type": labels.get(record_type, record_type), "occurred_at": occurred_at,
                "content": content if isinstance(content, dict) else {"text": cleaned(row.get("text"))},
                "product_ref": product, "enterprise_ref": enterprise, "location": location,
                "attachment_ids": json_cell(row.get("attachment_ids"), []), "provenance": {"imported": True, **({"source_timezone": source_timezone} if source_timezone else {})},
                "data_nature": row.get("data_nature") or "observed"}
