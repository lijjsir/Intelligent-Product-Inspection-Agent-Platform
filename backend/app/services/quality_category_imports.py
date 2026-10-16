"""Preview and atomically import administrator-supplied category definitions."""
import csv
from io import BytesIO, StringIO

from app.core.exceptions import ValidationError
from app.schemas.quality_risk import ProductCategoryCreate


def read_category_rows(filename, content):
    if not content or len(content) > 5 * 1024 * 1024:
        raise ValidationError("目录文件不能为空，且不能超过5MB")
    try:
        if filename.lower().endswith('.csv'):
            reader = csv.DictReader(StringIO(content.decode('utf-8-sig')))
            headers = reader.fieldnames or []
            if len(headers) != len(set(headers)):
                raise ValueError('duplicate headers')
            rows = []
            for row in reader:
                if None in row:
                    raise ValueError('extra columns')
                rows.append(row)
                if len(rows) > 1000:
                    raise ValueError('too many rows')
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                values = workbook.active.iter_rows(values_only=True)
                headers = next(values)
                if len(headers) != len(set(headers)):
                    raise ValueError('duplicate headers')
                rows = []
                for values_row in values:
                    if not any(value not in (None, '') for value in values_row):
                        continue
                    rows.append(dict(zip(headers, values_row)))
                    if len(rows) > 1000:
                        raise ValueError('too many rows')
            finally:
                workbook.close()
        else:
            raise ValueError('unsupported file')
    except Exception as exc:
        raise ValidationError("请上传有效UTF-8 CSV或XLSX目录，列名不可重复，最多1000行") from exc
    if not rows:
        raise ValidationError("目录中没有类别记录")
    return rows


async def preview_categories(svc, filename, content, common_basis=None):
    svc.require_role({'admin'})
    raw_rows = read_category_rows(filename, content)
    valid, errors = [], []
    for index, row in enumerate(raw_rows, 2):
        try:
            name = row.get('name') or row.get('类别名称') or row.get('产品类别')
            basis = row.get('classification_basis') or row.get('分类依据') or row.get('依据') or common_basis
            name, basis = svc.category_definition(name, basis)
            valid.append({'name': name, 'classification_basis': basis})
        except ValidationError as exc:
            errors.append({'row': index, 'message': exc.message})
    return {'rows': valid, 'errors': errors, 'total_count': len(raw_rows), 'valid_count': len(valid),
            'can_confirm': bool(valid) and not errors}


async def confirm_categories(svc, payload):
    svc.require_role({'admin'})
    from sqlalchemy import select
    from app.models.quality_risk import ProductCategory
    from app.models.organization import Organization
    created, reused, ids = 0, 0, []
    async with svc.db.begin_nested():
        await svc.db.scalar(select(Organization.id).where(Organization.id == svc.org_id).with_for_update())
        for definition in payload.rows:
            name, basis = svc.category_definition(definition.name, definition.classification_basis)
            previous = await svc.db.scalar(select(ProductCategory.id).where(ProductCategory.org_id == svc.org_id,
                ProductCategory.name == name, ProductCategory.deleted_at.is_(None)).with_for_update())
            row = await svc.create_category(ProductCategoryCreate(name=name, classification_basis=basis))
            ids.append(row['id'])
            if previous:
                reused += 1
            else:
                created += 1
    return {'created_count': created, 'reused_count': reused, 'category_ids': ids}
