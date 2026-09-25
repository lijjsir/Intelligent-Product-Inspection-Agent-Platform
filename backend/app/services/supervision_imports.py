"""Preview-first imports for region dictionaries and source-backed risk cases."""

from __future__ import annotations

import base64
import csv
import io

from app.core.exceptions import ForbiddenError, ValidationError
from app.core.ids import uuid7
from app.models.supervision import SupervisionRun
from app.models.supervision import SupervisionRecord
from sqlalchemy import select
from app.schemas.supervision import RecordCreate


def read_rows(filename, content):
    if len(content) > 10 * 1024 * 1024:
        raise ValidationError("文件不能超过10MB")
    try:
        if filename.lower().endswith(".csv"):
            rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            values = wb.active.iter_rows(values_only=True)
            headers = next(values)
            rows = [dict(zip(headers, row)) for row in values]
            wb.close()
        else:
            raise ValueError()
    except Exception as exc:
        raise ValidationError("请上传UTF-8 CSV或有效XLSX文件") from exc
    if not rows or len(rows) > 5000:
        raise ValidationError("文件应包含1–5000行记录")
    return rows


async def preview(svc, kind, filename, content, mapping):
    await svc.require_enabled()
    svc.require_role({"admin"} if kind == "regions" else {"user", "expert"})
    rows = read_rows(filename, content)
    file_codes = {str(row.get(mapping.get("code", "code")) or "").strip() for row in rows}
    normalized, errors, codes = [], [], set()
    for index, row in enumerate(rows, start=2):

        def value(key):
            result = row.get(mapping.get(key, key))
            return result if result not in (None, "") else None

        try:
            code, name = str(value("code") or "").strip(), str(value("name") or "").strip()
            if code in codes:
                raise ValueError("文件内编号重复")
            codes.add(code)
            if kind == "regions":
                parent_code = str(value("parent_code") or "").strip()
                parent_id = value("parent_id")
                if parent_code:
                    if parent_code == code:
                        raise ValueError("地区不能成为自己的上级")
                    existing_parent = await svc.db.scalar(
                        select(SupervisionRecord).where(
                            SupervisionRecord.org_id == svc.org_id,
                            SupervisionRecord.kind == "regions",
                            SupervisionRecord.code == parent_code,
                            SupervisionRecord.status == "active",
                        )
                    )
                    if existing_parent:
                        parent_id = existing_parent.id
                    elif parent_code not in file_codes:
                        raise ValueError("上级地区编码不在字典或本文件中")
                data = {
                    "dictionary_version": str(value("dictionary_version") or "2026"),
                    "parent_id": parent_id,
                }
            else:
                data = {
                    "description": value("description"),
                    "occurred_at": value("occurred_at"),
                    "product_category": value("product_category"),
                    "enterprise_id": value("enterprise_id"),
                    "product_sku_id": value("product_sku_id"),
                    "batch_id": value("batch_id"),
                    "inspection_standard_id": value("inspection_standard_id"),
                    "complaint_region_id": value("complaint_region_id"),
                    "production_region_id": value("production_region_id"),
                    "sampling_region_id": value("sampling_region_id"),
                    "evidence": [
                        {
                            "evidence_id": f"{code}-source",
                            "source_type": value("source_type") or "complaint",
                            "source_id": value("source_id"),
                            "occurred_at": value("occurred_at"),
                            "text": value("source_text") or value("description"),
                            "nature": "observed",
                        }
                    ],
                }
            payload = RecordCreate(code=code, name=name, data=data)
            payload.data = await svc.validate_data(kind, data)
            existing = await svc.db.scalar(
                select(SupervisionRecord.id).where(
                    SupervisionRecord.org_id == svc.org_id,
                    SupervisionRecord.kind == kind,
                    SupervisionRecord.code == code,
                )
            )
            if existing:
                raise ValueError("编号已存在")
            normalized.append(
                payload.model_dump(mode="json")
                | ({"parent_code": parent_code} if kind == "regions" else {})
            )
        except Exception as exc:
            errors.append({"row": index, "message": str(exc)[:500]})
    if kind == "regions":
        parents = {row["code"]: row.get("parent_code") for row in normalized}
        for code in parents:
            seen = set()
            current = code
            while current and current in parents:
                if current in seen:
                    errors.append({"row": code, "message": "地区层级存在循环"})
                    break
                seen.add(current)
                current = parents[current]
    run = SupervisionRun(
        org_id=svc.org_id,
        record_id=svc.actor,
        input_version=1,
        request_key=f"import:{uuid7()}",
        agent=f"import:{kind}",
        status="preview",
        iteration=0,
        created_by=svc.actor,
        snapshot={
            "kind": kind,
            "filename": filename,
            "base64": base64.b64encode(content).decode(),
            "mapping": mapping,
            "rows": normalized,
            "errors": errors,
        },
    )
    svc.db.add(run)
    await svc.db.flush()
    return {
        "preview_id": run.id,
        "valid_count": len(normalized),
        "rows": normalized[:100],
        "errors": errors,
        "can_confirm": bool(normalized) and not errors,
    }


async def confirm(svc, kind, preview_id):
    from sqlalchemy import select

    run = await svc.db.scalar(
        select(SupervisionRun)
        .where(
            SupervisionRun.id == preview_id,
            SupervisionRun.org_id == svc.org_id,
            SupervisionRun.created_by == svc.actor,
            SupervisionRun.agent == f"import:{kind}",
        )
        .with_for_update()
    )
    if not run:
        raise ForbiddenError("预览不存在或不属于当前账号")
    if run.status == "completed":
        return run.output
    if run.status != "preview" or run.snapshot["errors"]:
        raise ValidationError("请修正错误后重新预览")
    created = []
    async with svc.db.begin_nested():
        remaining = list(run.snapshot["rows"])
        created_codes = {}
        while remaining:
            ready = [
                row
                for row in remaining
                if not row.get("parent_code")
                or row["parent_code"] in created_codes
                or row["parent_code"] not in {r["code"] for r in remaining}
            ]
            if not ready:
                raise ValidationError("地区层级存在循环")
            for row in ready:
                payload = {k: v for k, v in row.items() if k != "parent_code"}
                if row.get("parent_code") in created_codes:
                    payload["data"] = {
                        **payload["data"],
                        "parent_id": created_codes[row["parent_code"]],
                    }
                result = await svc.create(kind, RecordCreate.model_validate(payload))
                created.append(result["id"])
                created_codes[row["code"]] = result["id"]
                remaining.remove(row)
        run.status = "completed"
        run.output = {"accepted": len(created), "record_ids": created}
    return run.output
