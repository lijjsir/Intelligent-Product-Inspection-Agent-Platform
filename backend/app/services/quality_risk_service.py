from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import String, func, select

from app.core.exceptions import ConflictError, ForbiddenError, NotFoundError, ValidationError
from app.core.ids import uuid7
from app.models.quality_risk import (
    PhysicalSample,
    ProductCategory,
    ProductIdentifier,
    QualityAttachment,
    QualityDataSource,
    QualityEvidenceItem,
    QualityIngestionJob,
    QualityProduct,
    QualityRiskAssessment,
    QualityRiskCase,
    QualitySourceRecord,
    RiskPolicy,
    StandardExecutionRule,
)
from app.schemas.quality_risk import QualityEventCreate
from app.services.object_storage import build_object_storage

BUSINESS_READ = {"admin", "user", "expert", "platform_operator"}
BUSINESS_WRITE = {"admin", "user", "expert"}


def stable_hash(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode()
    ).hexdigest()


def clean(value: Any) -> str:
    return str(value or "").strip()


def utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


class QualityRiskService:
    def __init__(self, db, current):
        self.db = db
        self.current = current
        self.org_id = str(current.org_id)
        self.actor = str(current.user_id)
        self.storage = build_object_storage()

    def require_role(self, allowed: set[str]) -> None:
        if self.current.role not in allowed:
            raise ForbiddenError("当前角色无权执行此操作")

    async def product_catalog(self, *, include_inactive: bool = False) -> dict:
        self.require_role(BUSINESS_READ | {"algorithm_engineer"})
        category_stmt = select(ProductCategory).where(ProductCategory.org_id == self.org_id, ProductCategory.deleted_at.is_(None))
        product_stmt = select(QualityProduct).where(QualityProduct.org_id == self.org_id, QualityProduct.deleted_at.is_(None))
        if not include_inactive:
            category_stmt = category_stmt.where(ProductCategory.is_active.is_(True))
            product_stmt = product_stmt.where(QualityProduct.is_active.is_(True))
        categories = list(await self.db.scalars(category_stmt.order_by(ProductCategory.code)))
        products = list(await self.db.scalars(product_stmt.order_by(QualityProduct.name)))
        category_map = {x.id: x.name for x in categories}
        identifiers = list(
            await self.db.scalars(
                select(ProductIdentifier).where(ProductIdentifier.org_id == self.org_id, ProductIdentifier.deleted_at.is_(None))
            )
        )
        by_product: dict[str, list[dict]] = {}
        for item in identifiers:
            by_product.setdefault(item.product_id, []).append(
                {
                    "id": item.id,
                    "identifier_type": item.identifier_type,
                    "identifier_value": item.identifier_value,
                    "source_id": item.source_id,
                }
            )
        return {
            "categories": [self.category_response(x) for x in categories],
            "products": [self.product_response(x, category_map.get(x.category_id), by_product.get(x.id, [])) for x in products],
        }

    async def create_category(self, payload) -> dict:
        self.require_role({"admin"})
        if await self.db.scalar(select(ProductCategory.id).where(ProductCategory.org_id == self.org_id, ProductCategory.code == payload.code, ProductCategory.deleted_at.is_(None))):
            raise ConflictError("产品类别编码已存在")
        if payload.parent_id:
            await self.category(payload.parent_id)
        row = ProductCategory(org_id=self.org_id, **payload.model_dump())
        self.db.add(row)
        await self.db.flush()
        return self.category_response(row)

    async def update_category(self, category_id: str, payload) -> dict:
        self.require_role({"admin"})
        row = await self.category(category_id)
        data = payload.model_dump(exclude_unset=True)
        if data.get("parent_id") == row.id:
            raise ValidationError("产品类别不能成为自己的上级")
        for key, value in data.items():
            setattr(row, key, value)
        await self.db.flush()
        return self.category_response(row)

    async def category(self, category_id: str) -> ProductCategory:
        row = await self.db.scalar(select(ProductCategory).where(ProductCategory.id == category_id, ProductCategory.org_id == self.org_id, ProductCategory.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("产品类别不存在")
        return row

    async def create_product(self, payload) -> dict:
        self.require_role({"admin"})
        category = await self.category(payload.category_id)
        data = payload.model_dump(exclude={"identifiers"})
        row = QualityProduct(org_id=self.org_id, **data)
        self.db.add(row)
        await self.db.flush()
        identifiers = await self.replace_identifiers(row.id, payload.identifiers)
        return self.product_response(row, category.name, identifiers)

    async def update_product(self, product_id: str, payload) -> dict:
        self.require_role({"admin"})
        row = await self.product(product_id)
        data = payload.model_dump(exclude_unset=True, exclude={"identifiers"})
        if data.get("category_id"):
            await self.category(data["category_id"])
        for key, value in data.items():
            setattr(row, key, value)
        identifiers = None
        if payload.identifiers is not None:
            identifiers = await self.replace_identifiers(row.id, payload.identifiers)
        if identifiers is None:
            identifiers = [
                {"id": x.id, "identifier_type": x.identifier_type, "identifier_value": x.identifier_value, "source_id": x.source_id}
                for x in await self.db.scalars(select(ProductIdentifier).where(ProductIdentifier.product_id == row.id, ProductIdentifier.org_id == self.org_id, ProductIdentifier.deleted_at.is_(None)))
            ]
        category = await self.category(row.category_id)
        await self.db.flush()
        return self.product_response(row, category.name, identifiers)

    async def product(self, product_id: str) -> QualityProduct:
        row = await self.db.scalar(select(QualityProduct).where(QualityProduct.id == product_id, QualityProduct.org_id == self.org_id, QualityProduct.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("产品不存在")
        return row

    async def replace_identifiers(self, product_id: str, identifiers) -> list[dict]:
        existing = list(await self.db.scalars(select(ProductIdentifier).where(ProductIdentifier.product_id == product_id, ProductIdentifier.org_id == self.org_id, ProductIdentifier.deleted_at.is_(None))))
        for row in existing:
            row.deleted_at = utcnow()
        result = []
        for item in identifiers:
            duplicate = await self.db.scalar(select(ProductIdentifier.id).where(ProductIdentifier.org_id == self.org_id, ProductIdentifier.identifier_type == item.identifier_type, ProductIdentifier.identifier_value == item.identifier_value, ProductIdentifier.source_id == item.source_id, ProductIdentifier.deleted_at.is_(None)))
            if duplicate:
                raise ConflictError("产品标识已被其他产品使用")
            row = ProductIdentifier(org_id=self.org_id, product_id=product_id, **item.model_dump())
            self.db.add(row)
            await self.db.flush()
            result.append({"id": row.id, "identifier_type": row.identifier_type, "identifier_value": row.identifier_value, "source_id": row.source_id})
        return result

    async def list_sources(self) -> list[dict]:
        self.require_role(BUSINESS_READ | {"algorithm_engineer", "app_developer"})
        rows = list(await self.db.scalars(select(QualityDataSource).where(QualityDataSource.org_id == self.org_id, QualityDataSource.deleted_at.is_(None)).order_by(QualityDataSource.created_at.desc())))
        return [self.source_response(x) for x in rows]

    async def create_source(self, payload) -> dict:
        self.require_role({"admin", "expert"})
        if await self.db.scalar(select(QualityDataSource.id).where(QualityDataSource.org_id == self.org_id, QualityDataSource.code == payload.code, QualityDataSource.deleted_at.is_(None))):
            raise ConflictError("数据来源编码已存在")
        row = QualityDataSource(org_id=self.org_id, created_by=self.actor, **payload.model_dump())
        self.db.add(row)
        await self.db.flush()
        return self.source_response(row)

    async def update_source(self, source_id: str, payload) -> dict:
        self.require_role({"admin", "expert"})
        row = await self.source(source_id)
        for key, value in payload.model_dump(exclude_unset=True).items():
            setattr(row, key, value)
        await self.db.flush()
        return self.source_response(row)

    async def source(self, source_id: str) -> QualityDataSource:
        row = await self.db.scalar(select(QualityDataSource).where(QualityDataSource.id == source_id, QualityDataSource.org_id == self.org_id, QualityDataSource.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("数据来源不存在")
        return row

    async def create_event(self, payload: QualityEventCreate) -> dict:
        self.require_role(BUSINESS_WRITE)
        source = await self.source(payload.source_id)
        if source.status != "active":
            raise ValidationError("数据来源未启用")
        if payload.product_ref.category_id:
            await self.category(payload.product_ref.category_id)
        if payload.product_ref.product_id:
            product = await self.product(payload.product_ref.product_id)
            if payload.product_ref.category_id and product.category_id != payload.product_ref.category_id:
                raise ValidationError("产品不属于所选产品类别")
        await self.validate_attachments(payload.attachment_ids)
        data = payload.model_dump(mode="json")
        content_hash = stable_hash({k: data[k] for k in data if k not in {"ingestion_job_id"}})
        existing = await self.db.scalar(select(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.source_id == payload.source_id, QualitySourceRecord.content_hash == content_hash))
        if existing:
            return self.record_response(existing)
        row = QualitySourceRecord(
            org_id=self.org_id,
            source_id=payload.source_id,
            ingestion_job_id=payload.ingestion_job_id,
            external_record_id=payload.external_record_id,
            record_type=payload.record_type,
            occurred_at=payload.occurred_at,
            content=payload.content,
            content_hash=content_hash,
            enterprise_ref=payload.enterprise_ref.model_dump(mode="json", exclude_none=True),
            product_ref=payload.product_ref.model_dump(mode="json", exclude_none=True),
            location=payload.location.model_dump(mode="json", exclude_none=True),
            attachment_ids=payload.attachment_ids,
            provenance=payload.provenance,
            authorization_scope=payload.authorization_scope,
            data_nature=payload.data_nature,
            normalization_status="normalized",
            normalization_errors=[],
            created_by=self.actor,
        )
        self.db.add(row)
        await self.db.flush()
        await self.materialize_evidence(row)
        return self.record_response(row)

    async def list_records(self, *, page: int, size: int, record_type: str | None, keyword: str | None) -> dict:
        self.require_role(BUSINESS_READ)
        stmt = select(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.deleted_at.is_(None))
        if record_type:
            stmt = stmt.where(QualitySourceRecord.record_type == record_type)
        if keyword:
            stmt = stmt.where(QualitySourceRecord.content.cast(String).contains(keyword, autoescape=True))
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery()))
        rows = list(await self.db.scalars(stmt.order_by(QualitySourceRecord.occurred_at.desc()).offset((page - 1) * size).limit(size)))
        return {"items": [self.record_response(x) for x in rows], "total": int(total or 0), "page": page, "size": size}

    async def get_record(self, record_id: str) -> dict:
        self.require_role(BUSINESS_READ)
        return self.record_response(await self.record(record_id))

    async def record(self, record_id: str) -> QualitySourceRecord:
        row = await self.db.scalar(select(QualitySourceRecord).where(QualitySourceRecord.id == record_id, QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("来源记录不存在")
        return row

    async def record_evidence(self, record_id: str) -> list[dict]:
        self.require_role(BUSINESS_READ)
        await self.record(record_id)
        rows = list(await self.db.scalars(select(QualityEvidenceItem).where(QualityEvidenceItem.org_id == self.org_id, QualityEvidenceItem.source_record_id == record_id, QualityEvidenceItem.deleted_at.is_(None)).order_by(QualityEvidenceItem.created_at)))
        return [self.evidence_response(x) for x in rows]

    async def materialize_evidence(self, record: QualitySourceRecord) -> None:
        text = clean(record.content.get("text") or record.content.get("description") or record.content.get("summary"))
        structured = {k: v for k, v in record.content.items() if k not in {"text", "description", "summary"}}
        items = [(None, text, structured)]
        items.extend((aid, None, {}) for aid in record.attachment_ids)
        for index, (attachment_id, evidence_text, evidence_data) in enumerate(items, 1):
            payload = {"record_id": record.id, "index": index, "text": evidence_text, "data": evidence_data, "attachment_id": attachment_id}
            row = QualityEvidenceItem(
                org_id=self.org_id,
                evidence_code=f"EV-{record.id[-8:].upper()}-{index:02d}",
                source_record_id=record.id,
                evidence_type="attachment" if attachment_id else record.record_type,
                text=evidence_text,
                structured_data=evidence_data,
                attachment_id=attachment_id,
                nature=record.data_nature,
                locator={"source_record_id": record.id, "external_record_id": record.external_record_id},
                content_hash=stable_hash(payload),
            )
            self.db.add(row)
        await self.db.flush()

    async def upload_attachment(self, file) -> dict:
        self.require_role(BUSINESS_WRITE)
        data = await file.read(50 * 1024 * 1024 + 1)
        if not data:
            raise ValidationError("附件不能为空")
        if len(data) > 50 * 1024 * 1024:
            raise ValidationError("单个附件不能超过50MB")
        name = Path(file.filename or "file.bin").name
        checksum = hashlib.sha256(data).hexdigest()
        object_key = f"quality-data/{self.org_id}/{checksum[:12]}-{name}"
        stored = self.storage.put_bytes(bucket="quality-data", object_key=object_key, data=data, content_type=file.content_type)
        row = QualityAttachment(org_id=self.org_id, file_name=name, mime_type=file.content_type or "application/octet-stream", size_bytes=len(data), sha256=checksum, bucket=str(stored.get("bucket") or "quality-data"), object_key=str(stored.get("object_key") or object_key), uploaded_by=self.actor)
        self.db.add(row)
        await self.db.flush()
        return self.attachment_response(row)

    async def validate_attachments(self, ids: list[str]) -> None:
        if not ids:
            return
        count = await self.db.scalar(select(func.count()).select_from(QualityAttachment).where(QualityAttachment.org_id == self.org_id, QualityAttachment.id.in_(ids), QualityAttachment.deleted_at.is_(None)))
        if int(count or 0) != len(set(ids)):
            raise ValidationError("存在不可用附件")

    async def preview_import(self, *, source_id: str, request_key: str, file, mapping: dict) -> dict:
        self.require_role(BUSINESS_WRITE)
        await self.source(source_id)
        existing = await self.db.scalar(select(QualityIngestionJob).where(QualityIngestionJob.org_id == self.org_id, QualityIngestionJob.request_key == request_key))
        if existing:
            return self.ingestion_response(existing)
        content = await file.read(20 * 1024 * 1024 + 1)
        if len(content) > 20 * 1024 * 1024:
            raise ValidationError("导入文件不能超过20MB")
        rows = self.parse_rows(file.filename or "", content)
        normalized, errors = [], []
        for index, row in enumerate(rows, 2):
            mapped = {target: row.get(source) for target, source in mapping.items()} if mapping else row
            try:
                normalized.append(self.row_to_event(mapped, source_id))
            except (TypeError, ValueError, ValidationError) as exc:
                errors.append({"row": index, "message": str(exc)})
        batch_code = await self.next_import_batch_code()
        job = QualityIngestionJob(org_id=self.org_id, source_id=source_id, request_key=request_key, import_batch_code=batch_code, status="preview", file_name=file.filename, mapping=mapping, raw_payload={"rows": normalized[:5000]}, total_count=len(rows), valid_count=len(normalized), error_count=len(errors), errors=errors[:1000], created_by=self.actor)
        self.db.add(job)
        await self.db.flush()
        return self.ingestion_response(job)

    async def confirm_import(self, job_id: str) -> dict:
        self.require_role(BUSINESS_WRITE)
        job = await self.ingestion_job(job_id, lock=True)
        if job.status == "confirmed":
            return self.ingestion_response(job)
        if job.status != "preview":
            raise ConflictError("导入任务当前不能确认")
        accepted = 0
        for raw in job.raw_payload.get("rows", []):
            payload = QualityEventCreate.model_validate({**raw, "ingestion_job_id": job.id, "provenance": {**raw.get("provenance", {}), "import_batch_code": job.import_batch_code}})
            await self.create_event(payload)
            accepted += 1
        job.status = "confirmed"
        job.valid_count = accepted
        await self.db.flush()
        return self.ingestion_response(job)

    async def cancel_import(self, job_id: str) -> dict:
        self.require_role(BUSINESS_WRITE)
        job = await self.ingestion_job(job_id, lock=True)
        if job.status == "confirmed":
            raise ConflictError("已确认导入不能取消")
        job.status = "cancelled"
        await self.db.flush()
        return self.ingestion_response(job)

    async def ingestion_job(self, job_id: str, *, lock: bool = False) -> QualityIngestionJob:
        stmt = select(QualityIngestionJob).where(QualityIngestionJob.id == job_id, QualityIngestionJob.org_id == self.org_id, QualityIngestionJob.deleted_at.is_(None))
        if lock:
            stmt = stmt.with_for_update()
        row = await self.db.scalar(stmt)
        if not row:
            raise NotFoundError("导入任务不存在")
        return row

    async def create_risk_case(self, payload) -> dict:
        self.require_role(BUSINESS_WRITE)
        record_count = await self.db.scalar(select(func.count()).select_from(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.id.in_(payload.source_record_ids)))
        if int(record_count or 0) != len(set(payload.source_record_ids)):
            raise ValidationError("风险案件包含不可用来源记录")
        evidence_ids = list(payload.evidence_ids)
        if not evidence_ids:
            evidence_ids = list(await self.db.scalars(select(QualityEvidenceItem.id).where(QualityEvidenceItem.org_id == self.org_id, QualityEvidenceItem.source_record_id.in_(payload.source_record_ids), QualityEvidenceItem.deleted_at.is_(None))))
        row = QualityRiskCase(org_id=self.org_id, code=f"RISK-{str(uuid7())[-12:].upper()}", title=payload.title, scope_type=payload.scope_type, scope=payload.scope, status="collecting", evidence_ids=evidence_ids, source_record_ids=payload.source_record_ids, assigned_to=payload.assigned_to, created_by=self.actor)
        self.db.add(row)
        await self.db.flush()
        return self.risk_case_response(row)

    async def list_risk_cases(self, *, page: int, size: int, status: str | None, keyword: str | None) -> dict:
        self.require_role(BUSINESS_READ)
        stmt = select(QualityRiskCase).where(QualityRiskCase.org_id == self.org_id, QualityRiskCase.deleted_at.is_(None))
        if status:
            stmt = stmt.where(QualityRiskCase.status == status)
        if keyword:
            stmt = stmt.where(QualityRiskCase.title.contains(keyword, autoescape=True))
        total = await self.db.scalar(select(func.count()).select_from(stmt.subquery()))
        rows = list(await self.db.scalars(stmt.order_by(QualityRiskCase.created_at.desc()).offset((page - 1) * size).limit(size)))
        return {"items": [self.risk_case_response(x) for x in rows], "total": int(total or 0), "page": page, "size": size}

    async def analyze_risk(self, risk_case_id: str, payload) -> dict:
        self.require_role({"user", "expert"})
        case = await self.risk_case(risk_case_id, lock=True)
        evidence = list(await self.db.scalars(select(QualityEvidenceItem).where(QualityEvidenceItem.org_id == self.org_id, QualityEvidenceItem.id.in_(case.evidence_ids), QualityEvidenceItem.deleted_at.is_(None))))
        observed = [x for x in evidence if x.nature == "observed"]
        records = list(await self.db.scalars(select(QualitySourceRecord).where(QualitySourceRecord.org_id == self.org_id, QualitySourceRecord.id.in_(case.source_record_ids))))
        policy = None
        if payload.policy_id:
            policy = await self.db.scalar(select(RiskPolicy).where(RiskPolicy.id == payload.policy_id, RiskPolicy.org_id == self.org_id, RiskPolicy.status == "published"))
            if not policy:
                raise ValidationError("风险政策不存在或尚未发布")
        explicit_levels = [clean(x.content.get("risk_level")) for x in records if clean(x.content.get("risk_level")) in {"low", "medium", "high", "critical"}]
        level = max(explicit_levels, key=lambda x: {"low": 1, "medium": 2, "high": 3, "critical": 4}[x]) if explicit_levels else "unknown"
        types = [clean(x.content.get("risk_type")) for x in records if clean(x.content.get("risk_type"))]
        causes = [clean(c) for x in records for c in (x.content.get("possible_causes") or []) if clean(c)]
        recommendations = [clean(c) for x in records for c in (x.content.get("recommendations") or []) if clean(c)]
        model_versions = {}
        opinion_records = [
            x
            for x in records
            if x.record_type in {"consumer_complaint", "public_opinion", "image", "video"}
        ]
        if opinion_records:
            try:
                from app.services.supervision_service import SupervisionService

                supervision = SupervisionService(self.db, self.current)
                model_call = await supervision.model_call(case.id)
                output = await supervision.execute_agent_via_manager(
                    {
                        "id": case.id,
                        "kind": "risk-case",
                        "version": 1,
                        "as_of": utcnow().isoformat(),
                        "knowledge_snapshot_id": stable_hash(
                            {"records": [x.content_hash for x in opinion_records]}
                        ),
                        "standard": None,
                        "data": {
                            "product_category": case.scope.get("name")
                            if case.scope_type == "category"
                            else case.scope.get("product_category"),
                            "description": "\n".join(
                                clean(x.content.get("text") or x.content.get("description"))
                                for x in opinion_records
                            ),
                            "occurred_at": max(x.occurred_at for x in opinion_records).isoformat(),
                            "evidence": [
                                {
                                    "evidence_id": x.id,
                                    "source_type": "image"
                                    if x.evidence_type == "attachment"
                                    else x.evidence_type,
                                    "source_id": x.source_record_id,
                                    "occurred_at": x.created_at.isoformat(),
                                    "text": x.text or "多媒体来源证据",
                                    "nature": x.nature,
                                }
                                for x in observed
                            ],
                        },
                    },
                    "public_opinion_monitoring",
                    workflow_run_id=case.id,
                    model_call=model_call,
                )
                agent_level = clean(output.get("risk_level"))
                if agent_level in {"low", "medium", "high", "critical"}:
                    level = agent_level
                signals = [clean(x) for x in output.get("normalized_signals", []) if clean(x)]
                if signals and not types:
                    types = signals
                causes.extend(clean(x) for x in output.get("possible_causes", []) if clean(x))
                recommendations.extend(
                    clean(x) for x in output.get("recommendations", []) if clean(x)
                )
                model_versions = dict(output.get("model_versions") or {})
            except Exception:  # noqa: BLE001 - model/provider failures must degrade to review
                model_versions = {"status": "unavailable_or_manual_review_required"}
        missing = [] if observed else ["真实来源证据"]
        if level == "unknown":
            missing.append("经规则或专家确认的风险等级")
        latest = await self.db.scalar(select(func.max(QualityRiskAssessment.version)).where(QualityRiskAssessment.org_id == self.org_id, QualityRiskAssessment.risk_case_id == case.id))
        assessment = QualityRiskAssessment(org_id=self.org_id, risk_case_id=case.id, version=int(latest or 0) + 1, risk_type=types[0] if types else None, risk_level=level, scope=case.scope, evidence_ids=[x.id for x in observed], standard_matches=[], conflicts=[], missing_inputs=missing, possible_causes=list(dict.fromkeys(causes)), recommendations=list(dict.fromkeys(recommendations)), trust_status="needs_review", probability=None, model_versions=model_versions, policy_version=policy.version if policy else None, knowledge_snapshot_id=stable_hash({"records": [(x.id, x.updated_at) for x in records], "evidence": [x.content_hash for x in evidence]}), status="awaiting_evidence" if missing else "awaiting_review", submitted_by=self.actor)
        self.db.add(assessment)
        case.status = assessment.status
        await self.db.flush()
        return self.assessment_response(assessment)

    async def review_risk(self, risk_case_id: str, payload) -> dict:
        self.require_role({"expert"})
        case = await self.risk_case(risk_case_id, lock=True)
        assessment = await self.db.scalar(select(QualityRiskAssessment).where(QualityRiskAssessment.org_id == self.org_id, QualityRiskAssessment.risk_case_id == case.id).order_by(QualityRiskAssessment.version.desc()).limit(1).with_for_update())
        if not assessment:
            raise ValidationError("请先生成风险研判草稿")
        if assessment.submitted_by == self.actor:
            raise ForbiddenError("不能复核本人提交的风险研判")
        if payload.decision == "accept":
            if assessment.risk_level == "unknown" and not payload.risk_level:
                raise ValidationError("通过研判前必须由专家确认风险等级")
            assessment.risk_level = payload.risk_level or assessment.risk_level
            assessment.risk_type = payload.risk_type or assessment.risk_type
            if payload.possible_causes is not None:
                assessment.possible_causes = payload.possible_causes
            if payload.recommendations is not None:
                assessment.recommendations = payload.recommendations
        assessment.reviewed_by = self.actor
        assessment.reviewed_at = utcnow()
        assessment.review_comment = payload.comment
        assessment.trust_status = "verified" if payload.decision == "accept" else "needs_evidence"
        assessment.status = "accepted" if payload.decision == "accept" else payload.decision
        case.status = "risk_assessed" if payload.decision == "accept" else "awaiting_evidence"
        await self.db.flush()
        return self.assessment_response(assessment)

    async def get_assessment(self, assessment_id: str) -> dict:
        self.require_role(BUSINESS_READ)
        row = await self.db.scalar(select(QualityRiskAssessment).where(QualityRiskAssessment.id == assessment_id, QualityRiskAssessment.org_id == self.org_id, QualityRiskAssessment.deleted_at.is_(None)))
        if not row:
            raise NotFoundError("风险研判不存在")
        return self.assessment_response(row)

    async def risk_case(self, risk_case_id: str, *, lock: bool = False) -> QualityRiskCase:
        stmt = select(QualityRiskCase).where(QualityRiskCase.id == risk_case_id, QualityRiskCase.org_id == self.org_id, QualityRiskCase.deleted_at.is_(None))
        if lock:
            stmt = stmt.with_for_update()
        row = await self.db.scalar(stmt)
        if not row:
            raise NotFoundError("风险案件不存在")
        return row

    async def list_policies(self) -> list[dict]:
        self.require_role(BUSINESS_READ)
        rows = list(await self.db.scalars(select(RiskPolicy).where(RiskPolicy.org_id == self.org_id, RiskPolicy.deleted_at.is_(None)).order_by(RiskPolicy.created_at.desc())))
        return [self.policy_response(x) for x in rows]

    async def create_policy(self, payload) -> dict:
        self.require_role({"admin", "expert"})
        row = RiskPolicy(org_id=self.org_id, code=payload.code, name=payload.name, version=payload.version, rules=payload.rules, status="draft", created_by=self.actor)
        self.db.add(row)
        await self.db.flush()
        return self.policy_response(row)

    async def publish_policy(self, policy_id: str) -> dict:
        self.require_role({"admin"})
        row = await self.db.scalar(select(RiskPolicy).where(RiskPolicy.id == policy_id, RiskPolicy.org_id == self.org_id, RiskPolicy.deleted_at.is_(None)).with_for_update())
        if not row:
            raise NotFoundError("风险政策不存在")
        row.status = "published"
        row.published_by = self.actor
        row.published_at = utcnow()
        await self.db.flush()
        return self.policy_response(row)

    async def list_standard_rules(self) -> list[dict]:
        self.require_role(BUSINESS_READ | {"app_developer"})
        rows = list(
            await self.db.scalars(
                select(StandardExecutionRule)
                .where(
                    StandardExecutionRule.org_id == self.org_id,
                    StandardExecutionRule.deleted_at.is_(None),
                )
                .order_by(StandardExecutionRule.created_at.desc())
            )
        )
        return [self.standard_rule_response(x) for x in rows]

    async def create_standard_rule(self, payload) -> dict:
        self.require_role({"admin", "expert"})
        from app.models.inspection_standard_library import InspectionStandardLibrary

        standard = await self.db.scalar(
            select(InspectionStandardLibrary).where(
                InspectionStandardLibrary.id == payload.standard_id,
                InspectionStandardLibrary.is_active.is_(True),
                InspectionStandardLibrary.deleted_at.is_(None),
                (InspectionStandardLibrary.org_id == self.org_id)
                | (InspectionStandardLibrary.org_id.is_(None)),
            )
        )
        if not standard:
            raise ValidationError("检测标准不存在或未启用")
        if payload.category_id:
            await self.category(payload.category_id)
        if payload.product_id:
            product = await self.product(payload.product_id)
            if payload.category_id and product.category_id != payload.category_id:
                raise ValidationError("规则产品不属于所选类别")
        row = StandardExecutionRule(
            org_id=self.org_id,
            created_by=self.actor,
            status="draft",
            **payload.model_dump(),
        )
        self.db.add(row)
        await self.db.flush()
        return self.standard_rule_response(row)

    async def publish_standard_rule(self, rule_id: str) -> dict:
        self.require_role({"admin"})
        row = await self.db.scalar(
            select(StandardExecutionRule)
            .where(
                StandardExecutionRule.id == rule_id,
                StandardExecutionRule.org_id == self.org_id,
                StandardExecutionRule.deleted_at.is_(None),
            )
            .with_for_update()
        )
        if not row:
            raise NotFoundError("标准执行规则不存在")
        if not row.clause_ref or not row.condition:
            raise ValidationError("规则缺少标准条款或执行条件")
        row.status = "published"
        row.reviewed_by = row.reviewed_by or self.actor
        row.published_by = self.actor
        row.published_at = utcnow()
        await self.db.flush()
        return self.standard_rule_response(row)

    async def list_physical_samples(self, sampling_plan_id: str | None = None) -> list[dict]:
        self.require_role(BUSINESS_READ)
        stmt = select(PhysicalSample).where(
            PhysicalSample.org_id == self.org_id,
            PhysicalSample.deleted_at.is_(None),
        )
        if sampling_plan_id:
            stmt = stmt.where(PhysicalSample.sampling_plan_id == sampling_plan_id)
        rows = list(await self.db.scalars(stmt.order_by(PhysicalSample.created_at.desc())))
        return [self.physical_sample_response(x) for x in rows]

    async def create_physical_sample(self, payload) -> dict:
        self.require_role({"user", "expert"})
        from app.models.organization import Organization
        from app.models.supervision import SupervisionRecord

        organization = await self.db.get(Organization, self.org_id)
        if not organization or not (organization.settings or {}).get(
            "laboratory_validation_enabled", False
        ):
            raise ForbiddenError("当前组织尚未启用实验室验证")
        plan = await self.db.scalar(
            select(SupervisionRecord).where(
                SupervisionRecord.id == payload.sampling_plan_id,
                SupervisionRecord.org_id == self.org_id,
                SupervisionRecord.kind == "sampling-plans",
                SupervisionRecord.status == "approved",
            )
        )
        if not plan:
            raise ValidationError("实物样品必须来自已批准监督抽查计划")
        if payload.product_id:
            await self.product(payload.product_id)
        row = PhysicalSample(
            org_id=self.org_id,
            sampling_plan_id=payload.sampling_plan_id,
            product_id=payload.product_id,
            production_batch_ref=payload.production_batch_ref,
            external_sample_code=payload.external_sample_code,
            unit_serial_ref=payload.unit_serial_ref,
            sampled_at=payload.sampled_at,
            sampling_location=payload.sampling_location.model_dump(
                mode="json", exclude_none=True
            ),
            custody_status="collected" if payload.sampled_at else "planned",
        )
        self.db.add(row)
        await self.db.flush()
        return self.physical_sample_response(row)

    def parse_rows(self, name: str, content: bytes) -> list[dict]:
        suffix = Path(name).suffix.lower()
        if suffix == ".csv":
            return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        if suffix in {".xlsx", ".xlsm"}:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            values = list(ws.iter_rows(values_only=True))
            if not values:
                return []
            headers = [clean(x) for x in values[0]]
            return [{headers[i]: row[i] for i in range(min(len(headers), len(row)))} for row in values[1:] if any(x is not None for x in row)]
        if suffix == ".json":
            value = json.loads(content.decode("utf-8-sig"))
            return value if isinstance(value, list) else [value]
        raise ValidationError("首版结构化导入仅支持CSV、XLSX和JSON；PDF、图片和视频请通过附件上传")

    def row_to_event(self, row: dict, source_id: str) -> dict:
        record_type = clean(row.get("record_type"))
        occurred_at = clean(row.get("occurred_at"))
        if not record_type or not occurred_at:
            raise ValidationError("缺少record_type或occurred_at")
        content = row.get("content")
        if isinstance(content, str):
            try:
                content = json.loads(content)
            except json.JSONDecodeError:
                content = {"text": content}
        return {
            "source_id": source_id,
            "external_record_id": clean(row.get("external_record_id")) or None,
            "record_type": record_type,
            "occurred_at": occurred_at,
            "content": content if isinstance(content, dict) else {"text": clean(row.get("text"))},
            "enterprise_ref": {},
            "product_ref": {"production_batch_ref": clean(row.get("production_batch_ref")) or None, "unit_serial_ref": clean(row.get("unit_serial_ref")) or None},
            "location": {"formatted_address": clean(row.get("formatted_address")) or None, "location_method": "imported"},
            "attachment_ids": [],
            "provenance": {"imported": True},
        }

    async def next_import_batch_code(self) -> str:
        day = utcnow().strftime("%Y%m%d")
        count = await self.db.scalar(select(func.count()).select_from(QualityIngestionJob).where(QualityIngestionJob.org_id == self.org_id, QualityIngestionJob.import_batch_code.like(f"IMPORT-%-{day}-%")))
        index = int(count or 0)
        suffix = ""
        while True:
            suffix = chr(ord("A") + index % 26) + suffix
            index = index // 26 - 1
            if index < 0:
                break
        return f"IMPORT-DATA-{day}-{suffix}"

    @staticmethod
    def category_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "code", "name", "parent_id", "description", "is_active", "created_at", "updated_at")}

    @staticmethod
    def product_response(x, category_name, identifiers) -> dict:
        return {**{k: getattr(x, k) for k in ("id", "org_id", "category_id", "name", "model", "brand", "manufacturer_enterprise_id", "attributes", "is_active", "created_at", "updated_at")}, "category_name": category_name, "identifiers": identifiers}

    @staticmethod
    def source_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "code", "name", "source_type", "connector_type", "config", "status", "created_by", "created_at", "updated_at")}

    @staticmethod
    def record_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "source_id", "ingestion_job_id", "external_record_id", "record_type", "occurred_at", "received_at", "content", "content_hash", "enterprise_ref", "product_ref", "location", "attachment_ids", "provenance", "authorization_scope", "data_nature", "normalization_status", "normalization_errors", "created_by", "created_at")}

    def attachment_response(self, x) -> dict:
        return {"id": x.id, "file_name": x.file_name, "mime_type": x.mime_type, "size_bytes": x.size_bytes, "sha256": x.sha256, "download_url": self.storage.presign_download_url(bucket=x.bucket, object_key=x.object_key)}

    @staticmethod
    def evidence_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "evidence_code", "source_record_id", "evidence_type", "text", "structured_data", "attachment_id", "nature", "locator", "content_hash", "created_at")}

    @staticmethod
    def ingestion_response(x) -> dict:
        return {"id": x.id, "source_id": x.source_id, "request_key": x.request_key, "import_batch_code": x.import_batch_code, "status": x.status, "file_name": x.file_name, "total_count": x.total_count, "valid_count": x.valid_count, "error_count": x.error_count, "errors": x.errors, "preview_rows": x.raw_payload.get("rows", [])[:100]}

    @staticmethod
    def risk_case_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "code", "title", "scope_type", "scope", "status", "evidence_ids", "source_record_ids", "assigned_to", "created_by", "created_at", "updated_at")}

    @staticmethod
    def assessment_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "risk_case_id", "version", "risk_type", "risk_level", "scope", "evidence_ids", "standard_matches", "conflicts", "missing_inputs", "possible_causes", "recommendations", "trust_status", "probability", "model_versions", "policy_version", "knowledge_snapshot_id", "status", "submitted_by", "reviewed_by", "reviewed_at", "review_comment", "created_at")}

    @staticmethod
    def policy_response(x) -> dict:
        return {k: getattr(x, k) for k in ("id", "org_id", "code", "name", "version", "rules", "status", "created_by", "published_by", "published_at", "created_at")}

    @staticmethod
    def standard_rule_response(x) -> dict:
        return {
            k: getattr(x, k)
            for k in (
                "id",
                "org_id",
                "code",
                "name",
                "version",
                "standard_id",
                "clause_ref",
                "page_ref",
                "category_id",
                "product_id",
                "indicator",
                "method",
                "unit",
                "condition",
                "decision_action",
                "review_policy",
                "status",
                "created_by",
                "reviewed_by",
                "published_by",
                "published_at",
                "created_at",
            )
        }

    @staticmethod
    def physical_sample_response(x) -> dict:
        return {
            k: getattr(x, k)
            for k in (
                "id",
                "org_id",
                "sampling_plan_id",
                "product_id",
                "production_batch_ref",
                "external_sample_code",
                "unit_serial_ref",
                "sampled_at",
                "sampling_location",
                "custody_status",
                "created_at",
            )
        }
